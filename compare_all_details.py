from __future__ import annotations

import argparse
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from compare_monthly_detail import build_markdown, infer_period_label, load_sheet_totals, safe_name


FOUR_MD_SECOPS_ONLY = {"Detail_Astra", "Detail_Daya", "Detail_MPM", "Detail_Wahana"}
FOUR_MD_RI_ONLY = {"Detail_Analytic", "Detail_Loyalty"}
FOUR_MD_RAW = {"Detail_Shared"}
NINE_MD_RI_AND_SECOPS = {
    "Detail_Anper",
    "Detail_Sinsen",
    "Detail_CDN",
    "Detail_MenaraAgung",
    "Detail_MSK",
    "Detail_HPM",
    "Detail_ITC",
    "Detail_TDM",
    "Detail_TRIO",
    "Detail_Zenobia",
}


def get_detail_sheets(path: Path) -> set[str]:
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    return {sheet for sheet in wb.sheetnames if sheet.startswith("Detail_")}


def get_filters(sheet_name: str, workbook_kind: str) -> tuple[bool, bool]:
    if workbook_kind == "4md":
        if sheet_name in FOUR_MD_SECOPS_ONLY:
            return False, True
        if sheet_name in FOUR_MD_RI_ONLY:
            return True, False
        if sheet_name in FOUR_MD_RAW:
            return False, False
    if workbook_kind == "9md" and sheet_name in NINE_MD_RI_AND_SECOPS:
        return True, True
    return False, False


def accumulate_pair(
    previous: Path,
    current: Path,
    group_by: str,
    workbook_kind: str,
) -> tuple[Decimal, Decimal, dict[str, Decimal], dict[str, Decimal], list[str]]:
    prev_sheets = get_detail_sheets(previous)
    curr_sheets = get_detail_sheets(current)
    common_sheets = sorted(prev_sheets & curr_sheets)

    previous_total = Decimal("0")
    current_total = Decimal("0")
    previous_grouped: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    current_grouped: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

    for sheet_name in common_sheets:
        exclude_ri, exclude_secops = get_filters(sheet_name, workbook_kind)
        prev_total, prev_map = load_sheet_totals(
            previous,
            sheet_name,
            group_by,
            exclude_reserved_instance=exclude_ri,
            exclude_secops=exclude_secops,
        )
        curr_total, curr_map = load_sheet_totals(
            current,
            sheet_name,
            group_by,
            exclude_reserved_instance=exclude_ri,
            exclude_secops=exclude_secops,
        )

        previous_total += prev_total
        current_total += curr_total

        for key, value in prev_map.items():
            previous_grouped[key] += value
        for key, value in curr_map.items():
            current_grouped[key] += value

    return previous_total, current_total, previous_grouped, current_grouped, common_sheets


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare all Detail_* sheets into one combined markdown report.")
    parser.add_argument("--previous-4md", type=Path)
    parser.add_argument("--current-4md", type=Path)
    parser.add_argument("--previous-9md", type=Path)
    parser.add_argument("--current-9md", type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--group-by", choices=["product", "meter", "resource"], default="product")
    parser.add_argument("--top", type=int, default=10)
    args = parser.parse_args()

    if not any([args.previous_4md and args.current_4md, args.previous_9md and args.current_9md]):
        parser.error("Provide at least one valid workbook pair: 4-MD or 9-MD.")

    previous_total = Decimal("0")
    current_total = Decimal("0")
    previous_grouped: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    current_grouped: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    included_sheets: list[str] = []

    if args.previous_4md and args.current_4md:
        p_total, c_total, p_map, c_map, sheets = accumulate_pair(
            args.previous_4md,
            args.current_4md,
            args.group_by,
            "4md",
        )
        previous_total += p_total
        current_total += c_total
        included_sheets.extend(sheets)
        for key, value in p_map.items():
            previous_grouped[key] += value
        for key, value in c_map.items():
            current_grouped[key] += value

    if args.previous_9md and args.current_9md:
        p_total, c_total, p_map, c_map, sheets = accumulate_pair(
            args.previous_9md,
            args.current_9md,
            args.group_by,
            "9md",
        )
        previous_total += p_total
        current_total += c_total
        included_sheets.extend(sheets)
        for key, value in p_map.items():
            previous_grouped[key] += value
        for key, value in c_map.items():
            current_grouped[key] += value

    changes = []
    all_keys = set(previous_grouped) | set(current_grouped)
    for key in all_keys:
        prev_value = previous_grouped.get(key, Decimal("0"))
        curr_value = current_grouped.get(key, Decimal("0"))
        delta = curr_value - prev_value
        changes.append((key, prev_value, curr_value, delta))

    increases = sorted((item for item in changes if item[3] > 0), key=lambda item: item[3], reverse=True)[: args.top]
    decreases = sorted((item for item in changes if item[3] < 0), key=lambda item: item[3])[: args.top]
    direction = "Naik" if current_total - previous_total > 0 else "Turun" if current_total - previous_total < 0 else "Flat"

    previous_source = args.previous_4md or args.previous_9md
    current_source = args.current_4md or args.current_9md
    previous_label = safe_name(infer_period_label(previous_source))
    current_label = safe_name(infer_period_label(current_source))
    output_path = args.output_dir / f"ALL_DETAILS-{previous_label}-vs-{current_label}.md"

    markdown = build_markdown(
        previous=previous_source,
        current=current_source,
        sheet="ALL_DETAILS",
        group_by=args.group_by,
        prev_total=previous_total,
        curr_total=current_total,
        direction=direction,
        increases=increases,
        decreases=decreases,
        top=args.top,
    )

    lines = markdown.splitlines()
    insert_at = 5 if len(lines) >= 5 else len(lines)
    lines.insert(insert_at, f"- Included sheets: `{', '.join(sorted(set(included_sheets)))}`")
    markdown = "\n".join(lines)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(markdown, encoding="utf-8")
    print(f"Markdown report saved: {output_path}")


if __name__ == "__main__":
    main()
