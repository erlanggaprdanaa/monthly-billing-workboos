from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


OUTPUT_PATH = Path(r"C:\Users\DYS\OneDrive\Dokumen\2026\template-summary.xlsx")
MAX_ROW_REF = 1000000
ACCOUNTING_IDR_FORMAT = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'
MONTH_FORMAT = "mmm-yy"


def payg_formula(sheet_name: str) -> str:
    return f'=SUM(INDIRECT("\'{sheet_name}\'!W2:W{MAX_ROW_REF}"))'


def ri_formula(sheet_name: str) -> str:
    return f'=SUM(INDIRECT("\'{sheet_name}\'!W2:W{MAX_ROW_REF}"))'


def apply_cell_style(cell, *, bold: bool = False, center: bool = False, fill: str | None = None) -> None:
    cell.font = Font(bold=bold)
    if center:
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if fill:
        cell.fill = PatternFill(fill_type="solid", fgColor=fill)
    thin = Side(style="thin", color="000000")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def build_summary_sheet() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 6
    ws.column_dimensions["I"].width = 18

    ws["A1"] = "Month:"
    ws["B1"] = None
    ws["A2"] = "Azure Usage"
    ws["A3"] = "No"
    ws["C3"] = "Astra"
    ws["D3"] = "Daya"
    ws["E3"] = "MPM"
    ws["F3"] = "Wahana"
    ws["G3"] = "Shared Apps"
    ws["G4"] = "Shared Analytics"
    ws["B4"] = "Payment Method"
    ws["B5"] = "PAYG"
    ws["B6"] = "Azure RI"
    ws["B7"] = "Total"

    ws["I3"] = (
        f'=SUM(INDIRECT("\'Detail_Shared\'!W2:W{MAX_ROW_REF}"),'
        f'INDIRECT("\'Detail_Loyalty\'!W2:W{MAX_ROW_REF}"))'
    )
    ws["I4"] = f'=SUM(INDIRECT("\'Detail_Analytic\'!W2:W{MAX_ROW_REF}"))'
    ws["I5"] = "=SUM(I3:I4)"
    ws["I6"] = f'=SUM(INDIRECT("\'Reserved Instance Shared\'!W2:W{MAX_ROW_REF}"))'
    ws["I7"] = "=SUM(I5:I6)"

    ws["C5"] = payg_formula("Detail_Astra")
    ws["D5"] = payg_formula("Detail_Daya")
    ws["E5"] = payg_formula("Detail_MPM")
    ws["F5"] = payg_formula("Detail_Wahana")

    ws["C6"] = ri_formula("Reserved Instance Astra")
    ws["D6"] = ri_formula("Reserved Instance Daya")
    ws["E6"] = ri_formula("Reserved Instance MPM")
    ws["F6"] = ri_formula("Reserved Instance Wahana")

    ws["C7"] = "=SUM(C5:C6)"
    ws["D7"] = "=SUM(D5:D6)"
    ws["E7"] = "=SUM(E5:E6)"
    ws["F7"] = "=SUM(F5:F6)"
    ws["G5"] = None
    ws["G6"] = None
    ws["G7"] = None
    ws["H3"] = "Rp"
    ws["H4"] = "Rp"
    ws["H5"] = "Rp"
    ws["H6"] = "Rp"
    ws["H7"] = "Rp"

    ws["A12"] = "Office 365 and Azure Support (Shared)"
    ws["A13"] = "Month:"
    ws["B13"] = "=B1"
    ws["A14"] = "No"
    ws["C14"] = "Item"
    ws["D14"] = "Qty"
    ws["E14"] = "Price"

    ws["A15"] = 1
    ws["B15"] = "Office"
    ws["C15"] = "O365"
    ws["D15"] = 1
    ws["E15"] = 32625

    ws["A16"] = 2
    ws["B16"] = "Support"
    ws["C16"] = "Azure Support"
    ws["D16"] = 1
    ws["E16"] = 38500000

    ws["B1"].number_format = MONTH_FORMAT
    ws["B13"].number_format = MONTH_FORMAT

    for cell_ref in ["I3", "I4", "I5", "I6", "I7", "C5", "D5", "E5", "F5", "C6", "D6", "E6", "F6", "C7", "D7", "E7", "F7", "E15", "E16"]:
        ws[cell_ref].number_format = ACCOUNTING_IDR_FORMAT

    header_fill = "D9EAF7"
    title_fill = "EDEDED"

    for ref in ["A2", "A12"]:
        apply_cell_style(ws[ref], bold=True, fill=title_fill)

    for ref in ["A3", "B4", "B5", "B6", "B7", "C3", "D3", "E3", "F3", "G3", "G4", "A14", "C14", "D14", "E14"]:
        apply_cell_style(ws[ref], bold=True, center=True, fill=header_fill)

    for ref in ["A1", "B1", "I3", "I4", "I5", "I6", "I7", "C5", "D5", "E5", "F5", "C6", "D6", "E6", "F6", "C7", "D7", "E7", "F7", "A13", "B13", "A15", "B15", "C15", "D15", "E15", "A16", "B16", "C16", "D16", "E16", "H3", "H4", "H5", "H6", "H7"]:
        apply_cell_style(ws[ref], center=True)

    for row in range(3, 8):
        ws[f"H{row}"] = "Rp"

    ws.freeze_panes = "A1"
    return wb


def main() -> None:
    wb = build_summary_sheet()
    wb.save(OUTPUT_PATH)
    print(f"Created template summary workbook: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
