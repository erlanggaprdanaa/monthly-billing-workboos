from __future__ import annotations

import argparse
import shutil
import tempfile
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_HEADERS = [
    "Period",
    "Agreement Name",
    "Company Name",
    "Tenant Primary Domain",
    "Publisher Subscription ID",
    "Publisher Subscription Name",
    "Product ID",
    "Product Name",
    "Category",
    "Subcategory",
    "Meter Name",
    "Usage Start Date",
    "Usage End Date",
    "Quantity",
    "Unit",
    "Currency",
    "Resource Group Name",
    "Resource Name",
    "Location",
    "Tags_JSON",
    "Tags_KeyValue",
    "Tags_Count",
    "Total",
]


def load_workbook_safe(path: Path):
    try:
        return load_workbook(path)
    except PermissionError:
        temp_copy = Path(tempfile.gettempdir()) / f"{path.stem}-copy{path.suffix}"
        shutil.copy2(path, temp_copy)
        return load_workbook(temp_copy)


def normalize(value: object) -> str:
    return "" if value is None else str(value).strip()


def get_headers(sheet) -> list[str]:
    return [normalize(cell.value) for cell in sheet[1]]


def main() -> None:
    parser = argparse.ArgumentParser(description="Append rows from Blank to Detail_Shared, then remove Blank sheet.")
    parser.add_argument("--workbook", required=True, type=Path)
    args = parser.parse_args()

    print(f"[INFO] Opening workbook : {args.workbook}")
    wb = load_workbook_safe(args.workbook)

    if "Blank" not in wb.sheetnames:
        print("[INFO] Blank sheet not found, skipping.")
        return
    if "Detail_Shared" not in wb.sheetnames:
        raise ValueError("Detail_Shared sheet not found.")

    blank_ws = wb["Blank"]
    shared_ws = wb["Detail_Shared"]

    if get_headers(blank_ws)[: len(EXPECTED_HEADERS)] != EXPECTED_HEADERS:
        raise ValueError("Blank sheet headers do not match expected shared detail headers.")
    if get_headers(shared_ws)[: len(EXPECTED_HEADERS)] != EXPECTED_HEADERS:
        raise ValueError("Detail_Shared headers do not match expected shared detail headers.")

    rows_to_append = []
    for row in blank_ws.iter_rows(min_row=2, values_only=True):
        trimmed = list(row[: len(EXPECTED_HEADERS)])
        if any(value is not None and value != "" for value in trimmed):
            rows_to_append.append(trimmed)

    if rows_to_append:
        print(f"[INFO] Appending {len(rows_to_append)} row(s) from Blank into Detail_Shared")
        for row in rows_to_append:
            shared_ws.append(row)
    else:
        print("[INFO] Blank sheet has no data rows, only deleting sheet.")

    wb.remove(blank_ws)
    print("[INFO] Blank sheet removed")
    wb.save(args.workbook)
    print("[DONE] Blank merged into Detail_Shared.")


if __name__ == "__main__":
    main()
