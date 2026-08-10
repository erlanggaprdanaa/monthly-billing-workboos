from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


SOURCE_HEADERS = [
    "Period",
    "Agreement Name",
    "Company Name",
    "Tenant Primary Domain",
    "Publisher Subscription ID",
    "Publisher Subscription Name",
    "Product ID",
    "Product Name",
    "Category",
    "Subcategory",
    "Meter Name",
    "Usage Start Date",
    "Usage End Date",
    "Quantity",
    "Unit",
    "Currency",
    "Resource Group Name",
    "Resource Name",
    "Location",
    "Tags_JSON",
    "Tags_KeyValue",
    "Tags_Count",
    "Total",
]

OVERRIDE_FROM_TARGET_HEADERS = {
    "Period",
    "Usage Start Date",
    "Usage End Date",
}


def normalize_header(value: object) -> str:
    return "" if value is None else str(value).strip()


def get_header_map(sheet) -> dict[str, int]:
    headers = {}
    for col_idx, cell in enumerate(sheet[1], start=1):
        header = normalize_header(cell.value)
        if header:
            headers[header] = col_idx
    return headers


def find_next_empty_row(sheet) -> int:
    for row_idx in range(sheet.max_row, 0, -1):
        if any(sheet.cell(row_idx, col_idx).value is not None for col_idx in range(1, sheet.max_column + 1)):
            return row_idx + 1
    return 1


def find_first_data_row(sheet) -> int:
    for row_idx in range(2, sheet.max_row + 1):
        if any(sheet.cell(row_idx, col_idx).value is not None for col_idx in range(1, sheet.max_column + 1)):
            return row_idx
    raise ValueError("Target sheet does not contain a template data row.")


def validate_headers(source_sheet, target_sheet) -> tuple[dict[str, int], dict[str, int]]:
    source_header_map = get_header_map(source_sheet)
    target_header_map = get_header_map(target_sheet)

    missing_source = [header for header in SOURCE_HEADERS if header not in source_header_map]
    missing_target = [header for header in SOURCE_HEADERS if header not in target_header_map]
    if missing_source or missing_target:
        raise ValueError(
            f"Header mismatch. Missing in source: {missing_source or '-'}; missing in target: {missing_target or '-'}"
        )

    return source_header_map, target_header_map


def collect_source_rows(source_sheet, source_header_map: dict[str, int]) -> list[dict[str, object]]:
    rows = []
    for source_row_idx in range(2, source_sheet.max_row + 1):
        row_data = {header: source_sheet.cell(source_row_idx, source_header_map[header]).value for header in SOURCE_HEADERS}
        if any(value is not None for value in row_data.values()):
            rows.append(row_data)
    return rows


def append_rows_to_target(target_sheet, source_rows: list[dict[str, object]], target_header_map: dict[str, int]) -> tuple[int, int]:
    start_row = find_next_empty_row(target_sheet)
    template_row = find_first_data_row(target_sheet)

    for copied, source_row in enumerate(source_rows):
        target_row_idx = start_row + copied
        for header in SOURCE_HEADERS:
            target_col_idx = target_header_map[header]
            if header in OVERRIDE_FROM_TARGET_HEADERS:
                target_sheet.cell(target_row_idx, target_col_idx).value = target_sheet.cell(template_row, target_col_idx).value
            else:
                target_sheet.cell(target_row_idx, target_col_idx).value = source_row[header]

    return start_row, len(source_rows)


def copy_rows(source_path: Path, source_sheet_name: str, target_path: Path, target_sheet_names: list[str]) -> list[tuple[str, int, int]]:
    source_wb = load_workbook(source_path, data_only=False)
    target_wb = load_workbook(target_path, data_only=False)

    source_sheet = source_wb[source_sheet_name]
    source_header_map = get_header_map(source_sheet)
    source_rows = collect_source_rows(source_sheet, source_header_map)
    if not source_rows:
        raise ValueError(f"No data rows found in source sheet '{source_sheet_name}'.")

    results = []
    for target_sheet_name in target_sheet_names:
        target_sheet = target_wb[target_sheet_name]
        _, target_header_map = validate_headers(source_sheet, target_sheet)
        start_row, copied = append_rows_to_target(target_sheet, source_rows, target_header_map)
        results.append((target_sheet_name, start_row, copied))

    target_wb.save(target_path)
    return results


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Append rows from a Google SecOps source sheet into a draft workbook detail sheet."
    )
    parser.add_argument("--source", required=True, type=Path, help="Path to source workbook.")
    parser.add_argument("--source-sheet", required=True, help="Source sheet name.")
    parser.add_argument("--target", required=True, type=Path, help="Path to target workbook.")
    parser.add_argument(
        "--target-sheets",
        required=True,
        help="Comma-separated target sheet names.",
    )
    args = parser.parse_args()

    target_sheet_names = [sheet.strip() for sheet in args.target_sheets.split(",") if sheet.strip()]
    if not target_sheet_names:
        raise ValueError("At least one target sheet must be provided.")

    results = copy_rows(args.source, args.source_sheet, args.target, target_sheet_names)
    for target_sheet_name, start_row, copied in results:
        print(f"Copied {copied} rows into '{target_sheet_name}' starting at row {start_row}.")


if __name__ == "__main__":
    main()
