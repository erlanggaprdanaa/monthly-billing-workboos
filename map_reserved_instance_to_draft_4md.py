from __future__ import annotations

import argparse
import shutil
import tempfile
from pathlib import Path
from typing import Any

import win32com.client as win32
from openpyxl import load_workbook


SOURCE_HEADERS = [
    "Period",
    "Agreement Name",
    "Company Name",
    "Tenant Primary Domain",
    "Publisher Subscription ID",
    "Publisher Subscription Name",
    "Product ID",
    "Product Name",
    "Category",
    "Subcategory",
    "Meter Name",
    "Usage Start Date",
    "Usage End Date",
    "Quantity",
    "Unit",
    "Currency",
    "Resource Group Name",
    "Resource Name",
    "Location",
    "Tags_JSON",
    "Tags_KeyValue",
    "Tags_Count",
    "Total",
    "MD Name",
]

DETAIL_HEADERS = SOURCE_HEADERS[:-1]

APPEND_RULES = {
    "Analytic": "Detail_Analytic",
    "Loyalty": "Detail_Loyalty",
}

ACCOUNTING_IDR_FORMAT = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'
ACCOUNTING_IDR_FORMAT_LOCAL = '_-[$Rp-421] * #.##0_-;\\-[$Rp-421] * #.##0_-;_-[$Rp-421] * "-"_-;_-@_-'
QUANTITY_NUMBER_FORMAT = "General"
XL_UP = -4162


def normalize_header(value: object) -> str:
    return "" if value is None else str(value).strip()


def load_source_workbook_safe(path: Path):
    try:
        return load_workbook(path, read_only=True, keep_links=False, data_only=False)
    except PermissionError:
        temp_copy = Path(tempfile.gettempdir()) / f"{path.stem}-copy{path.suffix}"
        shutil.copy2(path, temp_copy)
        return load_workbook(temp_copy, read_only=True, keep_links=False, data_only=False)


def get_openpyxl_header_map(sheet) -> dict[str, int]:
    headers: dict[str, int] = {}
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    for col_idx, value in enumerate(first_row, start=1):
        header = normalize_header(value)
        if header:
            headers[header] = col_idx
    return headers


def ensure_headers(header_map: dict[str, int], required_headers: list[str], sheet_name: str) -> None:
    missing = [header for header in required_headers if header not in header_map]
    if missing:
        raise ValueError(f"Sheet '{sheet_name}' is missing required headers: {missing}")


def collect_rows_by_md(source_sheet, source_header_map: dict[str, int]) -> dict[str, list[list[Any]]]:
    rows_by_md: dict[str, list[list[Any]]] = {}

    for row in source_sheet.iter_rows(min_row=2, values_only=True):
        row_values = [row[source_header_map[header] - 1] for header in SOURCE_HEADERS]
        if not any(value is not None and value != "" for value in row_values):
            continue

        md_name = row[source_header_map["MD Name"] - 1]
        if md_name is None or str(md_name).strip() == "":
            continue

        rows_by_md.setdefault(str(md_name).strip(), []).append(row_values)

    return rows_by_md


def get_excel_header_map(sheet) -> dict[str, int]:
    last_col = sheet.Cells(1, sheet.Columns.Count).End(-4159).Column
    values = sheet.Range(sheet.Cells(1, 1), sheet.Cells(1, last_col)).Value
    header_map: dict[str, int] = {}
    for col_idx in range(1, last_col + 1):
        value = values[0][col_idx - 1] if isinstance(values, tuple) else values
        header = normalize_header(value)
        if header:
            header_map[header] = col_idx
    return header_map


def to_excel_matrix(rows: list[list[Any]]) -> tuple[tuple[Any, ...], ...]:
    return tuple(tuple(row) for row in rows)


def apply_formats(sheet, start_row: int, end_row: int, include_md_name: bool) -> None:
    if end_row < start_row:
        return
    sheet.Range(sheet.Cells(start_row, 14), sheet.Cells(end_row, 14)).NumberFormat = QUANTITY_NUMBER_FORMAT
    sheet.Range(sheet.Cells(start_row, 15), sheet.Cells(end_row, 15)).NumberFormat = ACCOUNTING_IDR_FORMAT
    sheet.Range(sheet.Cells(start_row, 23), sheet.Cells(end_row, 23)).NumberFormat = ACCOUNTING_IDR_FORMAT
    try:
        sheet.Range(sheet.Cells(start_row, 14), sheet.Cells(end_row, 14)).NumberFormatLocal = QUANTITY_NUMBER_FORMAT
        sheet.Range(sheet.Cells(start_row, 15), sheet.Cells(end_row, 15)).NumberFormatLocal = ACCOUNTING_IDR_FORMAT_LOCAL
        sheet.Range(sheet.Cells(start_row, 23), sheet.Cells(end_row, 23)).NumberFormatLocal = ACCOUNTING_IDR_FORMAT_LOCAL
    except Exception:
        pass
    if include_md_name:
        sheet.Range(sheet.Cells(start_row, 24), sheet.Cells(end_row, 24)).NumberFormat = "General"


def worksheet_exists(workbook, name: str) -> bool:
    try:
        workbook.Worksheets(name)
        return True
    except Exception:
        return False


def delete_sheet_if_exists(workbook, name: str) -> None:
    if not worksheet_exists(workbook, name):
        return
    workbook.Worksheets(name).Delete()


def delete_reserved_instance_sheets(workbook) -> list[str]:
    sheet_names = [str(worksheet.Name) for worksheet in workbook.Worksheets]
    deleted: list[str] = []
    for sheet_name in sheet_names:
        if sheet_name.startswith("Reserved Instance "):
            workbook.Worksheets(sheet_name).Delete()
            deleted.append(sheet_name)
    return deleted


def delete_existing_reserved_instance_rows(sheet) -> int:
    header_map = get_excel_header_map(sheet)
    ensure_headers(header_map, DETAIL_HEADERS, sheet.Name)

    tags_json_col = header_map["Tags_JSON"]
    deleted = 0
    last_row = sheet.Cells(sheet.Rows.Count, 1).End(XL_UP).Row
    for row_idx in range(last_row, 1, -1):
        tags_json = normalize_header(sheet.Cells(row_idx, tags_json_col).Value)
        if tags_json == "Reserved Instance":
            sheet.Rows(row_idx).Delete()
            deleted += 1
    return deleted


def build_dynamic_new_sheet_rules(workbook, rows_by_md: dict[str, list[list[Any]]]) -> list[tuple[str, dict[str, str]]]:
    rules: list[tuple[str, dict[str, str]]] = []
    excluded_md = set(APPEND_RULES)

    for worksheet in workbook.Worksheets:
        sheet_name = str(worksheet.Name)
        if not sheet_name.startswith("Detail_"):
            continue

        md_name = sheet_name.removeprefix("Detail_")
        if md_name in excluded_md:
            continue
        if md_name not in rows_by_md:
            continue

        rules.append(
            (
                md_name,
                {
                    "sheet_name": f"Reserved Instance {md_name}",
                    "anchor_sheet": sheet_name,
                },
            )
        )

    return rules


def create_filtered_sheet(workbook, rows: list[list[Any]], sheet_name: str, anchor_sheet_name: str) -> tuple[str, int]:
    if not worksheet_exists(workbook, anchor_sheet_name):
        raise ValueError(f"Anchor sheet '{anchor_sheet_name}' not found in target workbook.")

    delete_sheet_if_exists(workbook, sheet_name)
    anchor_sheet = workbook.Worksheets(anchor_sheet_name)
    new_sheet = workbook.Worksheets.Add()
    new_sheet.Name = sheet_name
    new_sheet.Move(Before=anchor_sheet)

    header_matrix = (tuple(SOURCE_HEADERS),)
    new_sheet.Range(new_sheet.Cells(1, 1), new_sheet.Cells(1, len(SOURCE_HEADERS))).Value = header_matrix
    new_sheet.Range(new_sheet.Cells(1, 1), new_sheet.Cells(1, len(SOURCE_HEADERS))).Font.Bold = True

    if rows:
        data = to_excel_matrix(rows)
        end_row = 1 + len(rows)
        end_col = len(SOURCE_HEADERS)
        new_sheet.Range(new_sheet.Cells(2, 1), new_sheet.Cells(end_row, end_col)).Value = data
        apply_formats(new_sheet, 2, end_row, include_md_name=True)
        new_sheet.Range("A2").Select()
        new_sheet.Application.ActiveWindow.FreezePanes = True
        new_sheet.Range(new_sheet.Cells(1, 1), new_sheet.Cells(end_row, end_col)).AutoFilter()
        new_sheet.Columns.AutoFit()

    return sheet_name, len(rows)


def find_next_empty_row(sheet) -> int:
    last_row = sheet.Cells(sheet.Rows.Count, 1).End(XL_UP).Row
    if last_row == 1 and not any(sheet.Cells(1, col).Value for col in range(1, min(10, sheet.Columns.Count) + 1)):
        return 1
    return last_row + 1


def append_rows_to_detail_sheet(sheet, rows: list[list[Any]]) -> tuple[int, int]:
    header_map = get_excel_header_map(sheet)
    ensure_headers(header_map, DETAIL_HEADERS, sheet.Name)

    start_row = find_next_empty_row(sheet)
    for header in DETAIL_HEADERS:
        col_idx = header_map[header]
        values = [[row[SOURCE_HEADERS.index(header)]] for row in rows]
        end_row = start_row + len(rows) - 1
        sheet.Range(sheet.Cells(start_row, col_idx), sheet.Cells(end_row, col_idx)).Value = tuple(tuple(item) for item in values)

    if rows:
        apply_formats(sheet, start_row, start_row + len(rows) - 1, include_md_name=False)

    return start_row, len(rows)


def process(source_workbook_path: Path, target_workbook_path: Path, source_sheet_name: str) -> list[str]:
    print("[INFO] Loading source workbook")
    source_wb = load_source_workbook_safe(source_workbook_path)
    if source_sheet_name not in source_wb.sheetnames:
        raise ValueError(f"Source sheet '{source_sheet_name}' not found.")

    source_sheet = source_wb[source_sheet_name]
    source_header_map = get_openpyxl_header_map(source_sheet)
    ensure_headers(source_header_map, SOURCE_HEADERS, source_sheet.title)

    print("[INFO] Collecting rows by MD Name")
    rows_by_md = collect_rows_by_md(source_sheet, source_header_map)
    print(f"[INFO] Found MD groups: {', '.join(sorted(rows_by_md))}")

    print("[INFO] Opening target workbook in Excel")
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False

    workbook = None
    results: list[str] = []
    try:
        workbook = excel.Workbooks.Open(str(target_workbook_path))

        deleted_sheets = delete_reserved_instance_sheets(workbook)
        for sheet_name in deleted_sheets:
            results.append(f"[CLEAN] Deleted existing sheet '{sheet_name}'")

        dynamic_new_sheet_rules = build_dynamic_new_sheet_rules(workbook, rows_by_md)

        for md_name, rule in dynamic_new_sheet_rules:
            rows = rows_by_md.get(md_name, [])
            if not rows:
                results.append(f"[SKIP] {md_name}: no rows found in source workbook")
                continue
            print(f"[INFO] Creating sheet '{rule['sheet_name']}' with {len(rows)} row(s)")
            sheet_name, copied = create_filtered_sheet(workbook, rows, rule["sheet_name"], rule["anchor_sheet"])
            results.append(f"[NEW] {sheet_name}: copied {copied} row(s)")

        for md_name, target_sheet_name in APPEND_RULES.items():
            if not worksheet_exists(workbook, target_sheet_name):
                raise ValueError(f"Target sheet '{target_sheet_name}' not found in workbook.")

            target_sheet = workbook.Worksheets(target_sheet_name)
            deleted = delete_existing_reserved_instance_rows(target_sheet)
            if deleted:
                results.append(f"[CLEAN] {target_sheet_name}: deleted {deleted} existing Reserved Instance row(s)")

            rows = rows_by_md.get(md_name, [])
            if not rows:
                results.append(f"[SKIP] {md_name}: no rows found in source workbook")
                continue
            print(f"[INFO] Appending {len(rows)} row(s) into '{target_sheet_name}'")
            start_row, copied = append_rows_to_detail_sheet(target_sheet, rows)
            results.append(f"[APPEND] {target_sheet_name}: copied {copied} row(s) starting at row {start_row}")

        print("[INFO] Saving target workbook")
        workbook.Save()
        return results
    except Exception as exc:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
        excel.Quit()
        if isinstance(exc, PermissionError):
            raise PermissionError(f"Cannot save target workbook '{target_workbook_path}'. Please close the Excel file first.") from exc
        raise
    finally:
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=True)
            except Exception:
                pass
        excel.Quit()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Map Reserved Instance rows by MD Name into Draft-4MD workbook."
    )
    parser.add_argument("--source-workbook", required=True, type=Path, help="Path to reserved instance report workbook.")
    parser.add_argument("--target-workbook", required=True, type=Path, help="Path to Draft-4MD workbook.")
    parser.add_argument("--source-sheet", default="ReservedInstance", help="Source sheet name.")
    args = parser.parse_args()

    print(f"[INFO] Source workbook : {args.source_workbook}")
    print(f"[INFO] Target workbook : {args.target_workbook}")
    print(f"[INFO] Source sheet    : {args.source_sheet}")

    results = process(args.source_workbook, args.target_workbook, args.source_sheet)
    for line in results:
        print(line)
    print("[DONE] Reserved Instance mapping completed.")


if __name__ == "__main__":
    main()
