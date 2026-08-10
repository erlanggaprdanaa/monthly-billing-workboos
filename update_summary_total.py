from __future__ import annotations

import argparse
import re
import shutil
import tempfile
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


IDR_NUMBER_FORMAT = '"Rp" #,##0'
MONTH_FORMAT = "mmm-yy"

SUMMARY_COLUMNS = [
    "Astra",
    "Daya",
    "MPM",
    "Wahana",
    "Analytic",
    "Loyalty",
    "Shared",
    "Anper",
    "Sinsen",
    "CDN",
    "Menara Agung",
    "MSK",
    "HPM",
    "ITC",
    "TDM",
    "Trio",
    "Zenobia",
]

DETAIL_SHEET_TO_COLUMN = {
    "Detail_Astra": "Astra",
    "Detail_Daya": "Daya",
    "Detail_MPM": "MPM",
    "Detail_Wahana": "Wahana",
    "Detail_Analytic": "Analytic",
    "Detail_Loyalty": "Loyalty",
    "Detail_Shared": "Shared",
    "Detail_Anper": "Anper",
    "Detail_Sinsen": "Sinsen",
    "Detail_CDN": "CDN",
    "Detail_MenaraAgung": "Menara Agung",
    "Detail_MSK": "MSK",
    "Detail_HPM": "HPM",
    "Detail_ITC": "ITC",
    "Detail_TDM": "TDM",
    "Detail_TRIO": "Trio",
    "Detail_Zenobia": "Zenobia",
}

FRONT_RI_SHEET_TO_COLUMN = {
    "Reserved Instance Astra": "Astra",
    "Reserved Instance Daya": "Daya",
    "Reserved Instance MPM": "MPM",
    "Reserved Instance Wahana": "Wahana",
    "Reserved Instance Shared": "Shared",
}

FOUR_MD_SECOPS_PAYG_COLUMNS = {"Astra", "Daya", "MPM", "Wahana", "Shared"}
FOUR_MD_RI_PAYG_COLUMNS = {"Analytic", "Loyalty"}
NINE_MD_PAYG_COLUMNS = {"Anper", "Sinsen", "CDN", "Menara Agung", "MSK", "HPM", "ITC", "TDM", "Trio", "Zenobia"}

RI_PATTERN = re.compile(r"\bRI\b|Reserved Instance|Reservation", re.IGNORECASE)


def load_workbook_safe(path: Path, *, data_only: bool = False, read_only: bool = False):
    try:
        return load_workbook(path, data_only=data_only, keep_links=False, read_only=read_only)
    except PermissionError:
        temp_copy = Path(tempfile.gettempdir()) / f"{path.stem}-copy{path.suffix}"
        shutil.copy2(path, temp_copy)
        return load_workbook(temp_copy, data_only=data_only, keep_links=False, read_only=read_only)


def to_int(value: object) -> int:
    if value in (None, ""):
        return 0
    return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def round_decimal_total(value: Decimal) -> int:
    return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def header_index(headers: list[object], name: str) -> int:
    return headers.index(name) + 1


def sum_total_column(ws) -> int:
    headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    total_idx = headers.index("Total")
    total = Decimal("0")
    for row in ws.iter_rows(min_row=2, values_only=True):
        cell_value = row[total_idx]
        if cell_value not in (None, ""):
            total += Decimal(str(cell_value))
    return round_decimal_total(total)


def sum_filtered_total(ws, predicate) -> int:
    headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    total_idx = headers.index("Total")
    total = Decimal("0")
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_map = {str(headers[idx]): row[idx] for idx in range(len(headers))}
        if predicate(row_map):
            cell_value = row[total_idx]
            if cell_value not in (None, ""):
                total += Decimal(str(cell_value))
    return round_decimal_total(total)


def is_reserved_instance_row(row: dict[str, object]) -> bool:
    text = " ".join(
        str(row.get(key) or "")
        for key in ["Product Name", "Category", "Subcategory", "Meter Name", "Tags_JSON", "Tags_KeyValue"]
    )
    return bool(RI_PATTERN.search(text))


def is_google_secops_row(row: dict[str, object]) -> bool:
    resource_group = str(row.get("Resource Group Name") or "").lower()
    resource_name = str(row.get("Resource Name") or "").lower()
    product_name = str(row.get("Product Name") or "").lower()
    return "secops" in resource_group or "secops" in resource_name or "google secops" in product_name


def is_not_shared_license_row(row: dict[str, object]) -> bool:
    tags_keyvalue = str(row.get("Tags_KeyValue") or "").strip().lower()
    return tags_keyvalue != "shared allocation"


def collect_overall_totals(draft4: Path | None, draft9: Path | None) -> dict[str, int]:
    totals = {name: 0 for name in SUMMARY_COLUMNS}
    for workbook_path in [draft4, draft9]:
        if workbook_path is None:
            continue
        wb = load_workbook_safe(workbook_path, data_only=True, read_only=True)
        for sheet_name, column_name in DETAIL_SHEET_TO_COLUMN.items():
            if sheet_name in wb.sheetnames:
                totals[column_name] = sum_total_column(wb[sheet_name])
    return totals


def collect_reserved_instance_totals(draft4: Path | None, draft9: Path | None) -> dict[str, int]:
    totals = {name: 0 for name in SUMMARY_COLUMNS}

    if draft4 is not None:
        wb = load_workbook_safe(draft4, data_only=True, read_only=True)
        for sheet_name, column_name in FRONT_RI_SHEET_TO_COLUMN.items():
            if sheet_name in wb.sheetnames:
                if sheet_name == "Reserved Instance Shared":
                    totals[column_name] = sum_filtered_total(wb[sheet_name], is_not_shared_license_row)
                else:
                    totals[column_name] = sum_total_column(wb[sheet_name])

        for sheet_name, column_name in DETAIL_SHEET_TO_COLUMN.items():
            if column_name not in FOUR_MD_RI_PAYG_COLUMNS:
                continue
            if sheet_name in wb.sheetnames:
                totals[column_name] = sum_filtered_total(wb[sheet_name], is_reserved_instance_row)

    if draft9 is not None:
        wb = load_workbook_safe(draft9, data_only=True, read_only=True)
        for sheet_name, column_name in DETAIL_SHEET_TO_COLUMN.items():
            if column_name not in NINE_MD_PAYG_COLUMNS:
                continue
            if sheet_name in wb.sheetnames:
                totals[column_name] = sum_filtered_total(wb[sheet_name], is_reserved_instance_row)

    return totals


def collect_google_secops_totals(draft4: Path | None, draft9: Path | None) -> dict[str, int]:
    totals = {name: 0 for name in SUMMARY_COLUMNS}
    for workbook_path in [draft4, draft9]:
        if workbook_path is None:
            continue
        wb = load_workbook_safe(workbook_path, data_only=True, read_only=True)
        for sheet_name, column_name in DETAIL_SHEET_TO_COLUMN.items():
            if sheet_name not in wb.sheetnames:
                continue
            totals[column_name] = sum_filtered_total(wb[sheet_name], is_google_secops_row)
    return totals


def collect_shared_license_totals(template_workbook: Path) -> dict[str, int]:
    totals = {name: 0 for name in SUMMARY_COLUMNS}
    wb = load_workbook_safe(template_workbook, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    md_idx = headers.index("MD Name")
    total_idx = headers.index("Total")
    for row in ws.iter_rows(min_row=2, values_only=True):
        md_name = str(row[md_idx] or "").strip()
        if md_name == "Shared":
            totals["Shared"] += to_int(row[total_idx])
    return totals


def collect_reserved_instance_expected_total(ri_report_workbook: Path | None) -> int:
    if ri_report_workbook is None:
        return 0

    wb = load_workbook_safe(ri_report_workbook, data_only=True, read_only=True)
    if "Total 9-MD&4MD" not in wb.sheetnames:
        return 0

    ws = wb["Total 9-MD&4MD"]
    headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    md_name_idx = headers.index("MD Name")
    allocated_idx = headers.index("Allocated Total")
    total = Decimal("0")
    for row in ws.iter_rows(min_row=2, values_only=True):
        md_name = str(row[md_name_idx] or "").strip().upper()
        if md_name == "GRAND TOTAL":
            continue
        value = row[allocated_idx]
        if value not in (None, ""):
            total += Decimal(str(value))
    return round_decimal_total(total)


def ensure_month_sheet(workbook, sheet_name: str):
    if sheet_name in workbook.sheetnames:
        return workbook[sheet_name], False

    template_sheet = workbook[workbook.sheetnames[0]]
    new_sheet = workbook.copy_worksheet(template_sheet)
    new_sheet.title = sheet_name
    return new_sheet, True


def reset_sheet_layout(ws) -> None:
    ws["A1"] = None
    ws["A2"] = None
    for row_idx in range(2, 20):
        for col_idx in range(2, 25):
            ws.cell(row_idx, col_idx).value = None

    for col_idx, column_name in enumerate(SUMMARY_COLUMNS, start=2):
        ws.cell(2, col_idx).value = column_name
    ws.cell(2, len(SUMMARY_COLUMNS) + 2).value = "Total"

    labels = {
        3: "PAYG",
        4: "Reserved Instance",
        5: "Google SecOps",
        6: "Shared Licensi",
        8: "PAYG Expected",
        9: "PAYG Calculated",
        10: "Difference",
        11: "Status",
        13: "RI Expected",
        14: "RI Calculated",
        15: "Difference",
        16: "Status",
        18: "Last Updated",
        19: "Update Mode",
    }
    for row_idx, label in labels.items():
        ws.cell(row_idx, 1).value = label


def write_component_row(ws, row_idx: int, values_by_column: dict[str, int]) -> int:
    running_total = 0
    for col_offset, column_name in enumerate(SUMMARY_COLUMNS, start=2):
        value = int(values_by_column.get(column_name, 0))
        ws.cell(row_idx, col_offset).value = value
        ws.cell(row_idx, col_offset).number_format = IDR_NUMBER_FORMAT
        running_total += value

    total_col = len(SUMMARY_COLUMNS) + 2
    ws.cell(row_idx, total_col).value = running_total
    ws.cell(row_idx, total_col).number_format = IDR_NUMBER_FORMAT
    return running_total


def main() -> None:
    parser = argparse.ArgumentParser(description="Update summary-total workbook with monthly billing component totals.")
    parser.add_argument("--summary-workbook", required=True, type=Path)
    parser.add_argument("--draft-4md", type=Path)
    parser.add_argument("--draft-9md", type=Path)
    parser.add_argument("--reserved-instance-report-workbook", type=Path)
    parser.add_argument("--shared-license-workbook", required=True, type=Path)
    parser.add_argument("--sheet-name", required=True)
    parser.add_argument("--month-date", required=True, help="Month start date in YYYY-MM-DD")
    parser.add_argument("--payg-total-expected", required=True, type=int)
    args = parser.parse_args()

    if args.draft_4md is None and args.draft_9md is None:
        parser.error("At least one of --draft-4md or --draft-9md is required.")

    overall_totals = collect_overall_totals(args.draft_4md, args.draft_9md)
    ri_totals = collect_reserved_instance_totals(args.draft_4md, args.draft_9md)
    secops_totals = collect_google_secops_totals(args.draft_4md, args.draft_9md)
    shared_license_totals = collect_shared_license_totals(args.shared_license_workbook)
    ri_expected_total = collect_reserved_instance_expected_total(args.reserved_instance_report_workbook)

    payg_totals = {}
    for column_name in SUMMARY_COLUMNS:
        overall_value = overall_totals.get(column_name, 0)
        ri_value = ri_totals.get(column_name, 0)
        secops_value = secops_totals.get(column_name, 0)

        if column_name in {"Anper", "Sinsen", "CDN", "Menara Agung", "MSK", "HPM", "ITC", "TDM", "Trio", "Zenobia"}:
            payg_totals[column_name] = overall_value - ri_value - secops_value
        elif column_name in {"Astra", "Daya", "MPM", "Wahana"}:
            payg_totals[column_name] = overall_value - secops_value
        elif column_name in {"Analytic", "Loyalty"}:
            payg_totals[column_name] = overall_value - ri_value
        else:
            payg_totals[column_name] = overall_value

    print("[INFO] Component totals by MD")
    for column_name in SUMMARY_COLUMNS:
        print(
            "[INFO] "
            f"{column_name}: "
            f"Overall={overall_totals.get(column_name, 0):,} | "
            f"RI={ri_totals.get(column_name, 0):,} | "
            f"SecOps={secops_totals.get(column_name, 0):,} | "
            f"SharedLic={shared_license_totals.get(column_name, 0):,} | "
            f"PAYG={payg_totals.get(column_name, 0):,}"
        )

    wb = load_workbook_safe(args.summary_workbook, data_only=False)
    ws, created_new_sheet = ensure_month_sheet(wb, args.sheet_name)
    reset_sheet_layout(ws)

    ws["A1"] = datetime.fromisoformat(args.month_date)
    ws["A1"].number_format = MONTH_FORMAT

    payg_total = write_component_row(ws, 3, payg_totals)
    ri_calculated_total = write_component_row(ws, 4, ri_totals)
    write_component_row(ws, 5, secops_totals)
    write_component_row(ws, 6, shared_license_totals)

    ws["B8"] = args.payg_total_expected
    ws["B9"] = payg_total
    ws["B10"] = args.payg_total_expected - payg_total
    ws["B11"] = "MATCH" if args.payg_total_expected == payg_total else "MISMATCH"
    for cell_ref in ["B8", "B9", "B10"]:
        ws[cell_ref].number_format = IDR_NUMBER_FORMAT

    ws["B13"] = ri_expected_total
    ws["B14"] = ri_calculated_total
    ws["B15"] = ri_expected_total - ri_calculated_total
    ws["B16"] = "MATCH" if ri_expected_total == ri_calculated_total else "MISMATCH"
    for cell_ref in ["B13", "B14", "B15"]:
        ws[cell_ref].number_format = IDR_NUMBER_FORMAT

    ws["B18"] = datetime.now()
    ws["B18"].number_format = "yyyy-mm-dd hh:mm:ss"
    ws["B19"] = "Created new sheet" if created_new_sheet else "Refreshed existing sheet"

    wb.save(args.summary_workbook)
    print(f"[DONE] Updated summary workbook: {args.summary_workbook} | sheet: {args.sheet_name}")


if __name__ == "__main__":
    main()
