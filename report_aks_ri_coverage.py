from __future__ import annotations

import argparse
import json
import math
import subprocess
import urllib.parse
import urllib.request
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font


AZ_CLI = r"C:\Program Files (x86)\Microsoft SDKs\Azure\CLI2\wbin\az.cmd"
SUBSCRIPTION_ID = "a4cdaa2e-caa7-440f-b849-928a3ceb7e31"
RETAIL_PRICES_URL = "https://prices.azure.com/api/retail/prices"
USD_TO_IDR = 17955
MANUAL_USD_MONTHLY_PRICES = {
    ("Standard_B2s", "1 Year"): 22.50,
    ("Standard_DS2_v2", "1 Year"): 57.25,
}
DEFAULT_SHEETS = {
    "coverage": "Coverage",
    "recommendation": "Recommendation in Detail",
    "summary": "Summary Reserved",
}
DEFAULT_PAYG_PRICE_WORKBOOK = Path(r"H:\Work\Code-base\monthly-billing-workbooks\payg-price.xlsx")


def aks_name_from_rg(resource_group: str) -> str:
    parts = resource_group.split("_")
    if resource_group.startswith("MC_RG-") and len(parts) >= 4:
        name = parts[2]
    elif resource_group.startswith("MC_") and len(parts) >= 4:
        name = parts[2]
    else:
        return resource_group
    return name[:1].upper() + name[1:].lower() if name else resource_group


def rupiah(value: float) -> str:
    return f"Rp {value:,.0f}".replace(",", ".")


def decimal_comma(value: float, digits: int = 2) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def pct(value: float) -> str:
    return f"{value:.2f}%".replace(".", ",")


def round_half_up(value: float) -> int:
    return math.floor(value + 0.5)


def load_headers(ws: Any) -> dict[str, int]:
    headers = [cell for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    return {header: index for index, header in enumerate(headers)}


def to_number(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, int | float):
        return float(value)
    text = str(value).replace("Rp", "").replace(" ", "").replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def load_payg_prices(price_workbook: Path) -> dict[tuple[str, str], float]:
    wb = load_workbook(price_workbook, data_only=True, read_only=True)
    ws = wb["RI PAYG Price"] if "RI PAYG Price" in wb.sheetnames else wb.active
    idx = load_headers(ws)

    required = {"VM Size", "Region", "PAYG Monthly Price"}
    if not required.issubset(idx):
        missing = sorted(required - set(idx))
        raise ValueError(f"PAYG price workbook is missing required headers: {missing}")

    prices: dict[tuple[str, str], float] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        vm_size = str(row[idx["VM Size"]] or "").strip()
        region = str(row[idx["Region"]] or "").strip().lower()
        price = to_number(row[idx["PAYG Monthly Price"]])
        if vm_size and region and price > 0:
            prices[(vm_size, region)] = price
    return prices


def load_live_vmss() -> tuple[dict[str, dict[str, Any]], dict[tuple[str, str], int]]:
    try:
        completed = subprocess.run(
            [
                AZ_CLI,
                "vmss",
                "list",
                "--subscription",
                SUBSCRIPTION_ID,
                "--only-show-errors",
                "-o",
                "json",
            ],
            check=True,
            capture_output=True,
            text=True,
        )
    except subprocess.CalledProcessError as exc:
        print(f"[WARN] Azure VMSS live lookup failed; using workbook meter fallback. {exc}")
        return {}, {}
    items = json.loads(completed.stdout)
    by_name: dict[str, dict[str, Any]] = {}
    live_qty: dict[tuple[str, str], int] = defaultdict(int)
    for item in items:
        name = item.get("name") or ""
        if not name.startswith("aks-"):
            continue
        sku = (item.get("sku") or {}).get("name") or ""
        capacity = int((item.get("sku") or {}).get("capacity") or 0)
        resource_group = item.get("resourceGroup") or ""
        if not resource_group.startswith("MC_"):
            continue
        aks_name = aks_name_from_rg(resource_group)
        by_name[name] = {
            "aks_name": aks_name,
            "resource_group": resource_group,
            "sku": sku,
            "capacity": capacity,
        }
        if capacity > 0:
            live_qty[(aks_name, sku)] += capacity
    return by_name, live_qty


def vm_size_from_meter(meter_name: object, product_name: object) -> str:
    text = f"{meter_name or ''} {product_name or ''}"
    known_meter_map = {
        "D2 v3/D2s v3": "Standard_D2s_v3",
        "D4 v3/D4s v3": "Standard_D4s_v3",
        "F4s v2": "Standard_F4s_v2",
        "F8s v2": "Standard_F8s_v2",
        "B2s": "Standard_B2s",
        "DS2 v2": "Standard_DS2_v2",
    }
    for meter, vm_size in known_meter_map.items():
        if meter in text:
            return vm_size
    return ""


def read_reserved_details(
    reserved_report: Path,
) -> tuple[
    dict[tuple[str, str, str], dict[str, Any]],
    dict[str, float],
    set[str],
    dict[tuple[str, str], float],
]:
    wb = load_workbook(reserved_report, data_only=True, read_only=True)
    ws = wb["ReservedInstance"]
    idx = load_headers(ws)

    details: dict[tuple[str, str, str], dict[str, Any]] = defaultdict(
        lambda: {"qty": 0.0, "total": 0.0, "unit_price": 0.0, "resources": set()}
    )
    ri_total_by_aks: dict[str, float] = defaultdict(float)
    resource_groups: set[str] = set()
    unit_price_by_size_term: dict[tuple[str, str], float] = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        rg = str(row[idx["Resource Group Name"]] or "")
        subcategory = str(row[idx["Subcategory"]] or "")
        if "Reserved VM Instance" not in subcategory:
            continue

        parts = [part.strip() for part in subcategory.split(",")]
        if len(parts) < 4:
            continue

        vm_size = parts[1]
        term = parts[3]
        unit_price = float(row[idx["Unit"]] or 0)
        if unit_price > 0:
            unit_price_by_size_term.setdefault((vm_size, term), unit_price)

        if not rg.startswith("MC_"):
            continue

        aks_name = aks_name_from_rg(rg)
        key = (aks_name, vm_size, term)
        qty = float(row[idx["Quantity"]] or 0)
        total = float(row[idx["Total"]] or 0)
        resource_name = str(row[idx["Resource Name"]] or "")

        details[key]["qty"] += qty
        details[key]["total"] += total
        details[key]["unit_price"] = unit_price or details[key]["unit_price"]
        if resource_name:
            details[key]["resources"].add(resource_name)
        ri_total_by_aks[aks_name] += total
        resource_groups.add(rg)

    return details, ri_total_by_aks, resource_groups, unit_price_by_size_term


def read_payg(
    draft_workbooks: list[Path],
    resource_groups: set[str],
    live_by_resource: dict[str, dict[str, Any]],
) -> tuple[dict[str, float], dict[tuple[str, str], float]]:
    payg_by_aks: dict[str, float] = defaultdict(float)
    payg_by_aks_size: dict[tuple[str, str], float] = defaultdict(float)

    for workbook_path in draft_workbooks:
        wb = load_workbook(workbook_path, data_only=True, read_only=True)
        for sheet_name in wb.sheetnames:
            if not sheet_name.startswith("Detail_"):
                continue
            ws = wb[sheet_name]
            idx = load_headers(ws)
            required = {"Resource Group Name", "Category", "Tags_KeyValue", "Total", "Resource Name"}
            if not required.issubset(idx):
                continue
            for row in ws.iter_rows(min_row=2, values_only=True):
                rg = str(row[idx["Resource Group Name"]] or "")
                if rg not in resource_groups:
                    continue
                if str(row[idx["Category"]] or "") != "Virtual Machines":
                    continue
                if "Reserved Instance" in str(row[idx["Tags_KeyValue"]] or ""):
                    continue
                total = float(row[idx["Total"]] or 0)
                if total == 0:
                    continue

                resource_name = str(row[idx["Resource Name"]] or "")
                aks_name = aks_name_from_rg(rg)
                vm_size = live_by_resource.get(resource_name, {}).get("sku")
                if not vm_size and {"Meter Name", "Product Name"}.issubset(idx):
                    vm_size = vm_size_from_meter(row[idx["Meter Name"]], row[idx["Product Name"]])
                payg_by_aks[aks_name] += total
                if vm_size:
                    payg_by_aks_size[(aks_name, vm_size)] += total

    return payg_by_aks, payg_by_aks_size


def recommendation_note(estimated_qty: float, suggested_qty: int) -> str:
    if estimated_qty <= 0:
        return "No PAYG"
    if estimated_qty < 0.5:
        return "Combine with other cluster"
    if suggested_qty > estimated_qty:
        return "Round up"
    if suggested_qty < estimated_qty:
        return "Partial, round down"
    return "Exact"


def quantity_decision(estimated_qty: float, recommended_qty: int, coverage_qty: int) -> str:
    if estimated_qty <= 0:
        return "No PAYG"
    if estimated_qty < 0.5:
        return "Combine with other cluster"
    if recommended_qty == coverage_qty and estimated_qty != recommended_qty:
        return "Nearest qty, higher coverage"
    if recommended_qty < coverage_qty:
        return "Nearest qty, partial coverage"
    return "Exact"


def fetch_public_reservation_price(vm_size: str, region: str, term: str, currency: str = "IDR") -> float | None:
    term_months = 36 if term == "3 Years" else 12
    filters = [
        "serviceName eq 'Virtual Machines'",
        "priceType eq 'Reservation'",
        f"armRegionName eq '{region}'",
        f"reservationTerm eq '{term}'",
    ]
    query = {
        "currencyCode": currency,
        "$filter": " and ".join(filters),
    }
    url = f"{RETAIL_PRICES_URL}?{urllib.parse.urlencode(query)}"

    while url:
        with urllib.request.urlopen(url, timeout=30) as response:
            payload = json.loads(response.read().decode("utf-8"))

        prices = [
            float(item["unitPrice"])
            for item in payload.get("Items", [])
            if item.get("unitPrice") is not None
            and (
                item.get("armSkuName") == vm_size
                or item.get("skuName") == vm_size.replace("Standard_", "").replace("_", " ")
            )
        ]
        if prices:
            # Retail API reservation prices are full-term amounts; workbook unit prices are monthly.
            return min(prices) / term_months
        url = payload.get("NextPageLink")

    return None


def get_unit_price(
    vm_size: str,
    term: str,
    unit_price_by_size_term: dict[tuple[str, str], float],
    public_price_cache: dict[tuple[str, str], float | None],
) -> tuple[float, str]:
    workbook_price = unit_price_by_size_term.get((vm_size, term), 0.0)
    if workbook_price > 0:
        return workbook_price, "Workbook price"

    manual_usd_price = MANUAL_USD_MONTHLY_PRICES.get((vm_size, term))
    if manual_usd_price is not None:
        return manual_usd_price * USD_TO_IDR, "Manual USD price fallback"

    cache_key = (vm_size, term)
    if cache_key not in public_price_cache:
        public_price_cache[cache_key] = fetch_public_reservation_price(vm_size, "southeastasia", term)
    public_price = public_price_cache[cache_key]
    if public_price and public_price > 0:
        return public_price, "Public price fallback"

    return 0.0, "Need Price Input"


def build_coverage_rows(
    reserved_report: Path,
    draft_4md: Path,
    draft_9md: Path,
) -> list[dict[str, Any]]:
    _, ri_total_by_aks, resource_groups, _ = read_reserved_details(reserved_report)
    payg_by_aks, _ = read_payg([draft_4md, draft_9md], resource_groups, {})

    rows: list[dict[str, Any]] = []
    for aks_name in sorted(set(ri_total_by_aks) | set(payg_by_aks)):
        ri_total = ri_total_by_aks.get(aks_name, 0.0)
        payg_total = payg_by_aks.get(aks_name, 0.0)
        coverage = 1.0 if payg_total == 0 else (ri_total / (ri_total + payg_total) if (ri_total + payg_total) else 0.0)
        rows.append(
            {
                "aks_name": aks_name,
                "payg_left": payg_total,
                "coverage": coverage,
            }
        )

    rows.sort(key=lambda row: (row["coverage"], -row["payg_left"], row["aks_name"]))
    return rows


def write_coverage_excel(rows: list[dict[str, Any]], workbook_path: Path, sheet_name: str) -> None:
    wb = load_workbook(workbook_path)
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(sheet_name)

    ws.append(["AKS Name", "PAYG Left", "Coverage %"])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([row["aks_name"], row["payg_left"], row["coverage"]])

    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row_idx, 2).number_format = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'
        ws.cell(row_idx, 3).number_format = "0.00%"

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 16
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    wb.save(workbook_path)


def build_assessment_tables(
    reserved_report: Path,
    draft_4md: Path,
    draft_9md: Path,
    recommended_term: str,
    payg_price_workbook: Path,
    region: str,
    use_live_vmss: bool,
) -> dict[str, list[dict[str, Any]]]:
    live_by_resource, _ = load_live_vmss() if use_live_vmss else ({}, {})
    _, ri_total_by_aks, resource_groups, unit_price_by_size_term = read_reserved_details(reserved_report)
    payg_by_aks, payg_by_aks_size = read_payg([draft_4md, draft_9md], resource_groups, live_by_resource)
    public_price_cache: dict[tuple[str, str], float | None] = {}
    payg_prices = load_payg_prices(payg_price_workbook)

    coverage_rows: list[dict[str, Any]] = []
    for aks_name in sorted(set(ri_total_by_aks) | set(payg_by_aks)):
        ri_total = ri_total_by_aks.get(aks_name, 0.0)
        payg_total = payg_by_aks.get(aks_name, 0.0)
        coverage = 1.0 if payg_total == 0 else (ri_total / (ri_total + payg_total) if (ri_total + payg_total) else 0.0)
        coverage_rows.append({"AKS Name": aks_name, "PAYG Left": payg_total, "Coverage %": coverage})
    coverage_rows.sort(key=lambda row: (row["Coverage %"], -row["PAYG Left"], row["AKS Name"]))

    recommendation_rows: list[dict[str, Any]] = []
    recommendation_summary: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"payg_total": 0.0, "suggested_qty": 0, "combine_estimated_qty": 0.0, "clusters": []}
    )

    for (aks_name, vm_size), payg_total in sorted(payg_by_aks_size.items(), key=lambda item: (-item[1], item[0])):
        if payg_total <= 0:
            continue
        unit_price, price_source = get_unit_price(
            vm_size,
            recommended_term,
            unit_price_by_size_term,
            public_price_cache,
        )
        payg_monthly_price = payg_prices.get((vm_size, region.lower()), 0.0)
        if payg_monthly_price <= 0:
            estimated_qty = 0.0
            nearest_qty = 0
            coverage_qty = 0
            suggested_qty = 0
            estimated_ri_cost = 0.0
            potential_saving = 0.0
            recommended_ri = "Need Price Input"
            qty_decision = f"Need PAYG Price Input for {vm_size}"
            note = qty_decision
        else:
            estimated_qty = payg_total / payg_monthly_price
            nearest_qty = round_half_up(estimated_qty)
            coverage_qty = math.ceil(estimated_qty)
            suggested_qty = nearest_qty
            estimated_ri_cost = suggested_qty * unit_price if unit_price > 0 else 0.0
            potential_saving = payg_total - estimated_ri_cost if unit_price > 0 else 0.0
            recommended_ri = (
                f"{suggested_qty} x {vm_size} {recommended_term}"
                if suggested_qty > 0
                else "Combine with other cluster"
            )
            qty_decision = quantity_decision(estimated_qty, suggested_qty, coverage_qty)
            note = qty_decision
            if price_source == "Public price fallback":
                note = f"{note}; Public price fallback"
            elif price_source == "Manual USD price fallback":
                note = f"{note}; Manual USD price fallback"
            elif unit_price <= 0:
                note = f"{note}; Need {recommended_term} RI price input"

        summary = recommendation_summary[vm_size]
        summary["payg_total"] += payg_total
        summary["suggested_qty"] += suggested_qty
        if qty_decision == "Combine with other cluster":
            summary["combine_estimated_qty"] += estimated_qty
            summary["clusters"].append(f"{aks_name}({decimal_comma(estimated_qty)})")
        elif suggested_qty > 0:
            summary["clusters"].append(f"{aks_name}({suggested_qty})")

        recommendation_rows.append(
            {
                "AKS Name": aks_name,
                "VM Size PAYG": vm_size,
                "PAYG Total": payg_total,
                "PAYG Monthly Price": payg_monthly_price,
                "Recommended Term": recommended_term,
                "Estimated Workload Qty": estimated_qty,
                "Nearest Qty": nearest_qty,
                "Coverage Qty": coverage_qty,
                "Recommended Qty": suggested_qty,
                "Qty Decision": qty_decision,
                "RI Unit Price Acuan": unit_price,
                "Estimated RI Cost": estimated_ri_cost,
                "Potential Saving": potential_saving,
                "Recommended RI": recommended_ri,
                "Note": note,
            }
        )

    summary_rows: list[dict[str, Any]] = []
    for vm_size, summary in sorted(recommendation_summary.items()):
        combine_estimated_qty = float(summary["combine_estimated_qty"])
        combine_rounded_qty = round_half_up(combine_estimated_qty)
        suggested_qty = int(summary["suggested_qty"])
        summary_rows.append(
            {
                "VM Size PAYG": vm_size,
                "Total PAYG": float(summary["payg_total"]),
                "Suggested Purchase Qty": suggested_qty,
                "Combine Estimated Qty": combine_estimated_qty,
                "Combine Rounded Qty": combine_rounded_qty,
                "Total Planning Qty": suggested_qty + combine_rounded_qty,
                "Cluster": ", ".join(summary["clusters"]),
            }
        )

    return {
        "coverage": coverage_rows,
        "recommendation": recommendation_rows,
        "summary": summary_rows,
    }


def clear_or_create_sheet(wb: Any, sheet_name: str) -> Any:
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
        return ws
    return wb.create_sheet(sheet_name)


def write_rows(ws: Any, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def write_assessment_excel(
    tables: dict[str, list[dict[str, Any]]],
    workbook_path: Path,
    sheet_names: dict[str, str],
    selected_keys: set[str],
) -> dict[str, int]:
    wb = load_workbook(workbook_path)
    written: dict[str, int] = {}

    if "coverage" in selected_keys:
        ws = clear_or_create_sheet(wb, sheet_names["coverage"])
        headers = ["AKS Name", "PAYG Left", "Coverage %"]
        write_rows(ws, headers, tables["coverage"])
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row_idx, 2).number_format = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'
            ws.cell(row_idx, 3).number_format = "0.00%"
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 16
        written[sheet_names["coverage"]] = len(tables["coverage"])

    if "recommendation" in selected_keys:
        ws = clear_or_create_sheet(wb, sheet_names["recommendation"])
        headers = [
            "AKS Name",
            "VM Size PAYG",
            "PAYG Total",
            "PAYG Monthly Price",
            "Recommended Term",
            "Estimated Workload Qty",
            "Nearest Qty",
            "Coverage Qty",
            "Recommended Qty",
            "Qty Decision",
            "RI Unit Price Acuan",
            "Estimated RI Cost",
            "Potential Saving",
            "Recommended RI",
            "Note",
        ]
        write_rows(ws, headers, tables["recommendation"])
        for row_idx in range(2, ws.max_row + 1):
            for col_idx in (3, 4, 11, 12, 13):
                ws.cell(row_idx, col_idx).number_format = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'
            ws.cell(row_idx, 6).number_format = "0.00"
            for col_idx in (7, 8, 9):
                ws.cell(row_idx, col_idx).number_format = "0"
        widths = {
            "A": 18,
            "B": 20,
            "C": 18,
            "D": 20,
            "E": 16,
            "F": 24,
            "G": 14,
            "H": 14,
            "I": 18,
            "J": 28,
            "K": 20,
            "L": 18,
            "M": 18,
            "N": 30,
            "O": 44,
        }
        for column, width in widths.items():
            ws.column_dimensions[column].width = width
        written[sheet_names["recommendation"]] = len(tables["recommendation"])

    if "summary" in selected_keys:
        ws = clear_or_create_sheet(wb, sheet_names["summary"])
        headers = [
            "VM Size PAYG",
            "Total PAYG",
            "Suggested Purchase Qty",
            "Combine Estimated Qty",
            "Combine Rounded Qty",
            "Total Planning Qty",
            "Cluster",
        ]
        write_rows(ws, headers, tables["summary"])
        for row_idx in range(2, ws.max_row + 1):
            ws.cell(row_idx, 2).number_format = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'
            ws.cell(row_idx, 3).number_format = "0"
            ws.cell(row_idx, 4).number_format = "0.00"
            ws.cell(row_idx, 5).number_format = "0"
            ws.cell(row_idx, 6).number_format = "0"
        widths = {"A": 20, "B": 18, "C": 22, "D": 22, "E": 22, "F": 20, "G": 80}
        for column, width in widths.items():
            ws.column_dimensions[column].width = width
        written[sheet_names["summary"]] = len(tables["summary"])

    wb.save(workbook_path)
    return written


def markdown_from_assessment_tables(tables: dict[str, list[dict[str, Any]]], selected_keys: set[str]) -> str:
    sections: list[str] = ["# AKS RI Coverage", ""]
    if "coverage" in selected_keys:
        sections.extend(["## Coverage", "", "| AKS Name | PAYG Left | Coverage % |", "| --- | ---: | ---: |"])
        for row in tables["coverage"]:
            sections.append(f"| {row['AKS Name']} | {rupiah(float(row['PAYG Left']))} | {pct(float(row['Coverage %']) * 100)} |")
        sections.append("")
    if "recommendation" in selected_keys:
        sections.extend(
            [
                "## Recommendation in Detail",
                "",
                "| AKS Name | VM Size PAYG | PAYG Total | PAYG Monthly Price | Recommended Term | Estimated Workload Qty | Nearest Qty | Coverage Qty | Recommended Qty | Qty Decision | RI Unit Price Acuan | Estimated RI Cost | Potential Saving | Recommended RI | Note |",
                "| --- | --- | ---: | ---: | --- | ---: | ---: | ---: | ---: | --- | ---: | ---: | ---: | --- | --- |",
            ]
        )
        for row in tables["recommendation"]:
            sections.append(
                "| "
                + " | ".join(
                    [
                        row["AKS Name"],
                        row["VM Size PAYG"],
                        rupiah(float(row["PAYG Total"])),
                        rupiah(float(row["PAYG Monthly Price"])) if row["PAYG Monthly Price"] else "Need Price Input",
                        row["Recommended Term"],
                        decimal_comma(float(row["Estimated Workload Qty"])),
                        str(row["Nearest Qty"]),
                        str(row["Coverage Qty"]),
                        str(row["Recommended Qty"]),
                        row["Qty Decision"],
                        rupiah(float(row["RI Unit Price Acuan"])) if row["RI Unit Price Acuan"] else "Need Price Input",
                        rupiah(float(row["Estimated RI Cost"])) if row["Estimated RI Cost"] else "Need RI Price Input",
                        rupiah(float(row["Potential Saving"])),
                        row["Recommended RI"],
                        row["Note"],
                    ]
                )
                + " |"
            )
        sections.append("")
    if "summary" in selected_keys:
        sections.extend(
            [
                "## Summary Reserved",
                "",
                "| VM Size PAYG | Total PAYG | Suggested Purchase Qty | Combine Estimated Qty | Combine Rounded Qty | Total Planning Qty | Cluster |",
                "| --- | ---: | ---: | ---: | ---: | ---: | --- |",
            ]
        )
        for row in tables["summary"]:
            sections.append(
                "| "
                + " | ".join(
                    [
                        row["VM Size PAYG"],
                        rupiah(float(row["Total PAYG"])),
                        str(row["Suggested Purchase Qty"]),
                        decimal_comma(float(row["Combine Estimated Qty"])),
                        str(row["Combine Rounded Qty"]),
                        str(row["Total Planning Qty"]),
                        row["Cluster"],
                    ]
                )
                + " |"
            )
        sections.append("")
    return "\n".join(sections)


def resolve_selected_sheets(only_sheet: list[str] | None, ignore_sheet: list[str] | None) -> set[str]:
    aliases = {
        "coverage": "coverage",
        "recommendation": "recommendation",
        "recommendation in detail": "recommendation",
        "detail": "recommendation",
        "summary": "summary",
        "summary reserved": "summary",
    }

    def normalize_sheet(value: str) -> str:
        key = value.strip().lower()
        if key not in aliases:
            raise ValueError(f"Unknown sheet selector '{value}'. Use Coverage, Recommendation in Detail, or Summary Reserved.")
        return aliases[key]

    selected = {normalize_sheet(value) for value in only_sheet} if only_sheet else set(DEFAULT_SHEETS)
    ignored = {normalize_sheet(value) for value in ignore_sheet} if ignore_sheet else set()
    return selected - ignored


def build_report(
    reserved_report: Path,
    draft_4md: Path,
    draft_9md: Path,
    recommended_term: str,
) -> str:
    live_by_resource, live_qty = load_live_vmss()
    details, ri_total_by_aks, resource_groups, unit_price_by_size_term = read_reserved_details(reserved_report)
    payg_by_aks, payg_by_aks_size = read_payg([draft_4md, draft_9md], resource_groups, live_by_resource)
    public_price_cache: dict[tuple[str, str], float | None] = {}

    aks_names = sorted(set(ri_total_by_aks) | set(payg_by_aks))
    lines: list[str] = []
    lines.append("# AKS RI Coverage")
    lines.append("")
    lines.append("## AKS Summary")
    lines.append("")
    lines.append("| AKS Name | PAYG Left | Coverage % |")
    lines.append("| --- | ---: | ---: |")
    for aks_name in aks_names:
        ri_total = ri_total_by_aks.get(aks_name, 0.0)
        payg_total = payg_by_aks.get(aks_name, 0.0)
        coverage = 100.0 if payg_total == 0 else (ri_total / (ri_total + payg_total) * 100 if (ri_total + payg_total) else 0)
        lines.append(f"| {aks_name} | {rupiah(payg_total)} | {pct(coverage)} |")

    lines.append("")
    lines.append("## AKS Reservation Detail")
    lines.append("")
    lines.append(
        "| AKS Name | VM Size | Term | RI Qty | RI Total | Unit Price | Live Qty | Used By Resource Name | PAYG Left | Coverage % |"
    )
    lines.append("| --- | --- | --- | ---: | ---: | ---: | ---: | --- | ---: | ---: |")

    for key in sorted(details):
        aks_name, vm_size, term = key
        value = details[key]
        unit_price = float(value["unit_price"] or 0)
        payg_total = payg_by_aks.get(aks_name, 0.0)
        ri_total = ri_total_by_aks.get(aks_name, 0.0)
        payg_size = payg_by_aks_size.get((aks_name, vm_size), 0.0)
        coverage = 100.0 if payg_total == 0 else (ri_total / (ri_total + payg_total) * 100 if (ri_total + payg_total) else 0)
        live = live_qty.get((aks_name, vm_size), 0)

        resources = ", ".join(f"`{resource}`" for resource in sorted(value["resources"]))
        lines.append(
            "| "
            + " | ".join(
                [
                    aks_name,
                    vm_size,
                    term,
                    decimal_comma(float(value["qty"])),
                    rupiah(float(value["total"])),
                    rupiah(unit_price),
                    str(live),
                    resources,
                    rupiah(payg_total),
                    pct(coverage),
                ]
            )
            + " |"
        )

    lines.append("")
    lines.append("## AKS RI Recommendation")
    lines.append("")
    lines.append(
        "| AKS Name | VM Size PAYG | PAYG Total | Recommended Term | RI Unit Price Acuan | Estimated Qty | Suggested Purchase Qty | Recommended RI | Note |"
    )
    lines.append("| --- | --- | ---: | --- | ---: | ---: | ---: | --- | --- |")

    recommendation_summary: dict[str, dict[str, Any]] = defaultdict(
        lambda: {"payg_total": 0.0, "suggested_qty": 0, "combine_estimated_qty": 0.0, "clusters": []}
    )

    for (aks_name, vm_size), payg_total in sorted(payg_by_aks_size.items()):
        if payg_total <= 0:
            continue
        unit_price, price_source = get_unit_price(
            vm_size,
            recommended_term,
            unit_price_by_size_term,
            public_price_cache,
        )
        if unit_price <= 0:
            estimated_qty = 0.0
            suggested_qty = 0
            recommended_ri = "Need Price Input"
            note = f"No {recommended_term} RI price found"
            unit_price_text = "Need Price Input"
        else:
            estimated_qty = payg_total / unit_price
            suggested_qty = round_half_up(estimated_qty)
            recommended_ri = (
                f"{suggested_qty} x {vm_size} {recommended_term}"
                if suggested_qty > 0
                else "Combine with other cluster"
            )
            note = recommendation_note(estimated_qty, suggested_qty)
            if price_source == "Public price fallback":
                note = f"{note}; Public price fallback"
            elif price_source == "Manual USD price fallback":
                note = f"{note}; Manual USD price fallback"
            unit_price_text = rupiah(unit_price)

        summary = recommendation_summary[vm_size]
        summary["payg_total"] += payg_total
        summary["suggested_qty"] += suggested_qty
        if "Combine with other cluster" in note:
            summary["combine_estimated_qty"] += estimated_qty
            summary["clusters"].append(f"{aks_name}({decimal_comma(estimated_qty)})")
        elif suggested_qty > 0:
            summary["clusters"].append(f"{aks_name}({suggested_qty})")

        lines.append(
            "| "
            + " | ".join(
                [
                    aks_name,
                    vm_size,
                    rupiah(payg_total),
                    recommended_term,
                    unit_price_text,
                    decimal_comma(estimated_qty),
                    str(suggested_qty),
                    recommended_ri,
                    note,
                ]
            )
            + " |"
        )

    lines.append("")
    lines.append("## AKS RI Recommendation Summary")
    lines.append("")
    lines.append(
        "| VM Size PAYG | Total PAYG | Suggested Purchase Qty | Combine Estimated Qty | Combine Rounded Qty | Total Planning Qty | Cluster |"
    )
    lines.append("| --- | ---: | ---: | ---: | ---: | ---: | --- |")
    for vm_size, summary in sorted(recommendation_summary.items()):
        combine_estimated_qty = float(summary["combine_estimated_qty"])
        combine_rounded_qty = round_half_up(combine_estimated_qty)
        suggested_qty = int(summary["suggested_qty"])
        total_planning_qty = suggested_qty + combine_rounded_qty
        clusters = ", ".join(summary["clusters"])
        lines.append(
            "| "
            + " | ".join(
                [
                    vm_size,
                    rupiah(float(summary["payg_total"])),
                    str(suggested_qty),
                    decimal_comma(combine_estimated_qty),
                    str(combine_rounded_qty),
                    str(total_planning_qty),
                    clusters,
                ]
            )
            + " |"
        )

    return "\n".join(lines) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate AKS reserved instance coverage report.")
    parser.add_argument("--reserved-instance-report", required=True, type=Path)
    parser.add_argument("--draft-4md", required=True, type=Path)
    parser.add_argument("--draft-9md", required=True, type=Path)
    parser.add_argument("--output-md", type=Path)
    parser.add_argument("--recommended-term", default="1 Year", choices=["1 Year", "3 Years"])
    parser.add_argument("--payg-price-workbook", default=DEFAULT_PAYG_PRICE_WORKBOOK, type=Path)
    parser.add_argument("--region", default="Southeast Asia", help="PAYG price lookup region.")
    parser.add_argument("--use-live-vmss", action="store_true", help="Use Azure live VMSS lookup for PAYG VM size mapping.")
    parser.add_argument("--excel-output", type=Path, help="Optional Excel workbook output path.")
    parser.add_argument("--coverage-sheet", default="Coverage", help="Excel coverage sheet name.")
    parser.add_argument("--recommendation-sheet", default="Recommendation in Detail", help="Excel recommendation detail sheet name.")
    parser.add_argument("--summary-sheet", default="Summary Reserved", help="Excel summary reserved sheet name.")
    parser.add_argument(
        "--only-sheet",
        action="append",
        help="Generate only this sheet. Can be repeated. Values: Coverage, Recommendation in Detail, Summary Reserved.",
    )
    parser.add_argument(
        "--ignore-sheet",
        action="append",
        help="Skip this sheet. Can be repeated. Values: Coverage, Recommendation in Detail, Summary Reserved.",
    )
    args = parser.parse_args()

    if not args.output_md and not args.excel_output:
        parser.error("At least one output is required: --output-md or --excel-output")

    try:
        selected_sheets = resolve_selected_sheets(args.only_sheet, args.ignore_sheet)
    except ValueError as exc:
        parser.error(str(exc))

    if not selected_sheets:
        parser.error("No sheet selected after applying --only-sheet/--ignore-sheet.")

    tables = build_assessment_tables(
        args.reserved_instance_report,
        args.draft_4md,
        args.draft_9md,
        args.recommended_term,
        args.payg_price_workbook,
        args.region,
        args.use_live_vmss,
    )

    if args.excel_output:
        sheet_names = {
            "coverage": args.coverage_sheet,
            "recommendation": args.recommendation_sheet,
            "summary": args.summary_sheet,
        }
        written = write_assessment_excel(tables, args.excel_output, sheet_names, selected_sheets)
        for sheet_name, row_count in written.items():
            print(f"[DONE] Wrote {row_count} row(s) to {args.excel_output} sheet '{sheet_name}'")

    if args.output_md:
        args.output_md.parent.mkdir(parents=True, exist_ok=True)
        args.output_md.write_text(markdown_from_assessment_tables(tables, selected_sheets), encoding="utf-8")
        print(f"[DONE] AKS RI coverage assessment saved: {args.output_md}")


if __name__ == "__main__":
    main()
