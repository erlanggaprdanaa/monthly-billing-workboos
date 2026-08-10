from __future__ import annotations

import argparse
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


MAX_ROW_REF = 1_000_000
ACCOUNTING_IDR_FORMAT = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'
MONTH_FORMAT = "mmm-yy"

DETAIL_SHEETS = [
    "Detail_Astra",
    "Detail_Daya",
    "Detail_MPM",
    "Detail_Wahana",
    "Detail_Loyalty",
    "Detail_Analytic",
    "Detail_Shared",
]


def load_workbook_safe(path: Path):
    try:
        return load_workbook(path)
    except PermissionError:
        temp_copy = Path(tempfile.gettempdir()) / f"{path.stem}-copy{path.suffix}"
        shutil.copy2(path, temp_copy)
        return load_workbook(temp_copy)


def infer_month_start(workbook) -> datetime:
    for sheet_name in DETAIL_SHEETS:
        if sheet_name not in workbook.sheetnames:
            continue
        ws = workbook[sheet_name]
        headers = [cell.value for cell in ws[1]]
        if "Usage Start Date" not in headers:
            continue
        usage_idx = headers.index("Usage Start Date") + 1
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row_idx, usage_idx).value
            if isinstance(value, datetime):
                return datetime(value.year, value.month, 1)
    raise ValueError("Could not infer month from detail sheets.")


def payg_formula(sheet_name: str) -> str:
    return f"=SUM({sheet_name}!W2:W{MAX_ROW_REF})"


def ri_formula(sheet_name: str) -> str:
    return f"=SUM('Reserved Instance {sheet_name}'!W2:W{MAX_ROW_REF})"


def apply_cell_style(cell, *, bold: bool = False, center: bool = False, fill: str | None = None) -> None:
    cell.font = Font(bold=bold)
    if center:
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if fill:
        cell.fill = PatternFill(fill_type="solid", fgColor=fill)
    thin = Side(style="thin", color="000000")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def build_summary_sheet(ws, month_start: datetime) -> None:
    ws["A1"] = "Month:"
    ws["B1"] = month_start
    ws["A2"] = "Azure Usage"
    ws["A3"] = "No"
    ws["C3"] = "Astra"
    ws["D3"] = "Daya"
    ws["E3"] = "MPM"
    ws["F3"] = "Wahana"
    ws["G3"] = "Shared Apps"
    ws["G4"] = "Shared Analytics"
    ws["B4"] = "Payment Method"
    ws["B5"] = "PAYG"
    ws["B6"] = "Azure RI"
    ws["B7"] = "Total"

    ws["I3"] = f"=SUM(Detail_Shared!W2:W{MAX_ROW_REF},Detail_Loyalty!W2:W{MAX_ROW_REF})"
    ws["I4"] = f"=SUM(Detail_Analytic!W2:W{MAX_ROW_REF})"
    ws["I5"] = "=SUM(I3:I4)"
    ws["I6"] = f"=SUM('Reserved Instance Shared'!W2:W{MAX_ROW_REF})"
    ws["I7"] = "=SUM(I5:I6)"

    ws["C5"] = payg_formula("Detail_Astra")
    ws["D5"] = payg_formula("Detail_Daya")
    ws["E5"] = payg_formula("Detail_MPM")
    ws["F5"] = payg_formula("Detail_Wahana")

    ws["C6"] = ri_formula("Astra")
    ws["D6"] = ri_formula("Daya")
    ws["E6"] = ri_formula("MPM")
    ws["F6"] = ri_formula("Wahana")

    ws["C7"] = "=SUM(C5:C6)"
    ws["D7"] = "=SUM(D5:D6)"
    ws["E7"] = "=SUM(E5:E6)"
    ws["F7"] = "=SUM(F5:F6)"

    for row in range(3, 8):
        ws[f"H{row}"] = "Rp"

    ws["A12"] = "Office 365 and Azure Support (Shared)"
    ws["A13"] = "Month:"
    ws["B13"] = "=B1"
    ws["A14"] = "No"
    ws["C14"] = "Item"
    ws["D14"] = "Qty"
    ws["E14"] = "Price"

    ws["A15"] = 1
    ws["B15"] = "Office"
    ws["C15"] = "O365"
    ws["D15"] = 1
    ws["E15"] = 32625

    ws["A16"] = 2
    ws["B16"] = "Support"
    ws["C16"] = "Azure Support"
    ws["D16"] = 1
    ws["E16"] = 38500000

    ws["B1"].number_format = MONTH_FORMAT
    ws["B13"].number_format = MONTH_FORMAT

    for cell_ref in [
        "I3", "I4", "I5", "I6", "I7",
        "C5", "D5", "E5", "F5",
        "C6", "D6", "E6", "F6",
        "C7", "D7", "E7", "F7",
        "E15", "E16",
    ]:
        ws[cell_ref].number_format = ACCOUNTING_IDR_FORMAT

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 6
    ws.column_dimensions["I"].width = 18

    header_fill = "D9EAF7"
    title_fill = "EDEDED"

    for ref in ["A2", "A12"]:
        apply_cell_style(ws[ref], bold=True, fill=title_fill)

    for ref in ["A3", "B4", "B5", "B6", "B7", "C3", "D3", "E3", "F3", "G3", "G4", "A14", "C14", "D14", "E14"]:
        apply_cell_style(ws[ref], bold=True, center=True, fill=header_fill)

    for ref in [
        "A1", "B1", "I3", "I4", "I5", "I6", "I7",
        "C5", "D5", "E5", "F5", "C6", "D6", "E6", "F6",
        "C7", "D7", "E7", "F7", "A13", "B13",
        "A15", "B15", "C15", "D15", "E15",
        "A16", "B16", "C16", "D16", "E16",
        "H3", "H4", "H5", "H6", "H7",
    ]:
        apply_cell_style(ws[ref], center=True)


def create_summary(target_workbook: Path) -> None:
    print(f"[INFO] Opening target workbook : {target_workbook}")
    workbook = load_workbook_safe(target_workbook)
    month_start = infer_month_start(workbook)

    if "Summary" in workbook.sheetnames:
        workbook.remove(workbook["Summary"])

    summary_ws = workbook.create_sheet("Summary", 0)
    build_summary_sheet(summary_ws, month_start)

    print("[INFO] Saving workbook")
    workbook.save(target_workbook)
    print("[DONE] Summary sheet created.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Create Summary sheet directly inside Draft-4MD workbook.")
    parser.add_argument("--source", required=True, type=Path, help="Target Draft-4MD workbook path.")
    args = parser.parse_args()

    create_summary(args.source)


if __name__ == "__main__":
    main()
