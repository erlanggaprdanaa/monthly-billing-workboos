from __future__ import annotations

import argparse
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font


def normalize(value: object) -> str:
    return "" if value is None else str(value).strip()


def to_number(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, int | float):
        return float(value)
    text = str(value).replace("Rp", "").replace(" ", "").replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def round_half_up(value: float) -> int:
    return int(value + 0.5)


def load_headers(ws: Any) -> dict[str, int]:
    headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    return {normalize(header): index for index, header in enumerate(headers) if normalize(header)}


def aks_name_from_rg(resource_group: str) -> str:
    parts = resource_group.split("_")
    if resource_group.upper().startswith("MC_") and len(parts) >= 4:
        name = parts[2]
        return name[:1].upper() + name[1:].lower()
    return resource_group


def vm_size_from_subcategory(subcategory: object) -> str:
    parts = [part.strip() for part in normalize(subcategory).split(",")]
    if len(parts) >= 2 and parts[1].startswith("Standard_"):
        return parts[1]
    return ""


def vm_size_from_meter(meter_name: object, product_name: object) -> str:
    text = f"{normalize(meter_name)} {normalize(product_name)}"
    known_meter_map = {
        "D2 v3/D2s v3": "Standard_D2s_v3",
        "D4 v3/D4s v3": "Standard_D4s_v3",
        "F4s v2": "Standard_F4s_v2",
        "F8s v2": "Standard_F8s_v2",
        "B2s": "Standard_B2s",
        "DS2 v2": "Standard_DS2_v2",
    }
    for meter, vm_size in known_meter_map.items():
        if meter in text:
            return vm_size
    return ""


def load_payg_prices(price_workbook: Path) -> dict[tuple[str, str], float]:
    wb = load_workbook(price_workbook, data_only=True, read_only=True)
    ws = wb["RI PAYG Price"] if "RI PAYG Price" in wb.sheetnames else wb.active
    idx = load_headers(ws)

    required = ["VM Size", "Region", "PAYG Monthly Price"]
    missing = [header for header in required if header not in idx]
    if missing:
        raise ValueError(f"PAYG price workbook is missing required headers: {missing}")

    prices: dict[tuple[str, str], float] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        vm_size = normalize(row[idx["VM Size"]])
        region = normalize(row[idx["Region"]]).lower()
        price = to_number(row[idx["PAYG Monthly Price"]])
        if vm_size and region and price > 0:
            prices[(vm_size, region)] = price
    return prices


def collect_payg_totals(draft_workbooks: list[Path]) -> dict[tuple[str, str], float]:
    payg_totals: dict[tuple[str, str], float] = defaultdict(float)
    for workbook_path in draft_workbooks:
        wb = load_workbook(workbook_path, data_only=True, read_only=True)
        for ws in wb.worksheets:
            if not ws.title.startswith("Detail_"):
                continue
            idx = load_headers(ws)
            required = ["Resource Group Name", "Category", "Tags_JSON", "Meter Name", "Product Name", "Total"]
            if any(header not in idx for header in required):
                continue
            for row in ws.iter_rows(min_row=2, values_only=True):
                resource_group = normalize(row[idx["Resource Group Name"]])
                if not resource_group.upper().startswith("MC_"):
                    continue
                if normalize(row[idx["Category"]]) != "Virtual Machines":
                    continue
                if normalize(row[idx["Tags_JSON"]]) == "Reserved Instance":
                    continue
                total = to_number(row[idx["Total"]])
                if total <= 0:
                    continue
                vm_size = vm_size_from_meter(row[idx["Meter Name"]], row[idx["Product Name"]])
                if vm_size:
                    payg_totals[(aks_name_from_rg(resource_group), vm_size)] += total
    return payg_totals


def collect_reserved_quantities(reserved_report_workbook: Path) -> dict[tuple[str, str], float]:
    wb = load_workbook(reserved_report_workbook, data_only=True, read_only=True)
    ws = wb["ReservedInstance"]
    idx = load_headers(ws)

    required = ["Resource Group Name", "Subcategory", "Quantity"]
    missing = [header for header in required if header not in idx]
    if missing:
        raise ValueError(f"Reserved report workbook is missing required headers: {missing}")

    reserved_qty: dict[tuple[str, str], float] = defaultdict(float)
    for row in ws.iter_rows(min_row=2, values_only=True):
        resource_group = normalize(row[idx["Resource Group Name"]])
        if not resource_group.upper().startswith("MC_"):
            continue
        subcategory = normalize(row[idx["Subcategory"]])
        if "Reserved VM Instance" not in subcategory:
            continue
        vm_size = vm_size_from_subcategory(subcategory)
        if vm_size:
            reserved_qty[(aks_name_from_rg(resource_group), vm_size)] += to_number(row[idx["Quantity"]])
    return reserved_qty


def build_rows(args: argparse.Namespace) -> list[dict[str, Any]]:
    prices = load_payg_prices(args.payg_price_workbook)
    payg_totals = collect_payg_totals(args.draft_workbook)
    reserved_qty = collect_reserved_quantities(args.reserved_report_workbook)

    rows: list[dict[str, Any]] = []
    for cluster_name, vm_size in sorted(set(payg_totals) | set(reserved_qty)):
        payg_total = payg_totals.get((cluster_name, vm_size), 0.0)
        payg_price = prices.get((vm_size, args.region.lower()), 0.0)
        has_missing_price = payg_total > 0 and payg_price <= 0
        qty_payg = payg_total / payg_price if payg_price else 0.0
        qty_reserved = reserved_qty.get((cluster_name, vm_size), 0.0)
        notes = "OK"
        if has_missing_price:
            notes = f"Need PAYG Price Input for {vm_size}"
        rows.append(
            {
                "Cluster Name": cluster_name,
                "VM Size": vm_size,
                "QTY PAYG": None if has_missing_price else qty_payg,
                "QTY Reserved Instance": qty_reserved,
                "Total QTY": (0.0 if has_missing_price else qty_payg) + qty_reserved,
                "Total QTY Rounded": round_half_up((0.0 if has_missing_price else qty_payg) + qty_reserved),
                "Notes": notes,
            }
        )
    return rows


def markdown_table(rows: list[dict[str, Any]]) -> str:
    lines = [
        "| Cluster Name | VM Size | QTY PAYG | QTY Reserved Instance | Total QTY | Total QTY Rounded | Notes |",
        "| --- | --- | ---: | ---: | ---: | ---: | --- |",
    ]
    for row in rows:
        qty_payg = "" if row["QTY PAYG"] is None else f"{row['QTY PAYG']:.2f}".replace(".", ",")
        lines.append(
            "| "
            + " | ".join(
                [
                    row["Cluster Name"],
                    row["VM Size"],
                    qty_payg,
                    f"{row['QTY Reserved Instance']:.2f}".replace(".", ","),
                    f"{row['Total QTY']:.2f}".replace(".", ","),
                    str(row["Total QTY Rounded"]),
                    row["Notes"],
                ]
            )
            + " |"
        )
    return "\n".join(lines) + "\n"


def write_excel(rows: list[dict[str, Any]], workbook_path: Path, sheet_name: str) -> None:
    wb = load_workbook(workbook_path)
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(sheet_name)

    headers = ["Cluster Name", "VM Size", "QTY PAYG", "QTY Reserved Instance", "Total QTY", "Total QTY Rounded", "Notes"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([row[header] for header in headers])

    for row_idx in range(2, ws.max_row + 1):
        for col_idx in (3, 4, 5):
            ws.cell(row_idx, col_idx).number_format = "0.00"
        ws.cell(row_idx, 6).number_format = "0"

    widths = {"A": 22, "B": 20, "C": 14, "D": 24, "E": 14, "F": 18, "G": 36}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    wb.save(workbook_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate AKS cluster PAYG + Reserved Instance quantity summary.")
    parser.add_argument("--reserved-report-workbook", required=True, type=Path, help="Clean reserved-instance report workbook.")
    parser.add_argument("--payg-price-workbook", required=True, type=Path, help="PAYG price workbook.")
    parser.add_argument("--draft-workbook", required=True, type=Path, action="append", help="Draft workbook. Repeat for 4MD and 9MD.")
    parser.add_argument("--region", default="Southeast Asia", help="PAYG price lookup region.")
    parser.add_argument("--excel-output", type=Path, help="Optional Excel workbook output path.")
    parser.add_argument("--excel-sheet", default="Cluster QTY Summary", help="Excel output sheet name.")
    parser.add_argument("--output-md", type=Path, help="Optional markdown output path.")
    args = parser.parse_args()

    if not args.excel_output and not args.output_md:
        parser.error("At least one output is required: --excel-output or --output-md")

    rows = build_rows(args)
    if args.excel_output:
        write_excel(rows, args.excel_output, args.excel_sheet)
        print(f"[DONE] Wrote {len(rows)} row(s) to {args.excel_output} sheet '{args.excel_sheet}'")
    if args.output_md:
        args.output_md.parent.mkdir(parents=True, exist_ok=True)
        args.output_md.write_text(markdown_table(rows), encoding="utf-8")
        print(f"[DONE] Wrote {len(rows)} row(s) to {args.output_md}")


if __name__ == "__main__":
    main()
