from __future__ import annotations

import argparse
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font


def normalize(value: object) -> str:
    return "" if value is None else str(value).strip()


def to_number(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, int | float):
        return float(value)
    text = str(value).replace("Rp", "").replace(" ", "").replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalize(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def load_headers(ws: Any) -> dict[str, int]:
    headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    return {normalize(header): index for index, header in enumerate(headers) if normalize(header)}


def rupiah(value: float) -> str:
    return f"Rp {value:,.0f}".replace(",", ".")


def decimal_comma(value: float, digits: int = 2) -> str:
    return f"{value:.{digits}f}".replace(".", ",")


def percent(value: float) -> str:
    return f"{value:.2f}%".replace(".", ",")


def expiry_date_list(values: set[date]) -> str:
    return ", ".join(value.isoformat() for value in sorted(values))


def aks_name_from_rg(resource_group: str) -> str:
    parts = resource_group.split("_")
    if resource_group.upper().startswith("MC_") and len(parts) >= 4:
        name = parts[2]
        return name[:1].upper() + name[1:].lower()
    return resource_group


def vm_size_from_text(*parts: object) -> str:
    text = " ".join(normalize(part) for part in parts)
    match = re.search(r"Standard_([A-Za-z0-9]+s?_v\d+)", text, re.IGNORECASE)
    if match:
        return f"Standard_{match.group(1)}"

    known_meter_map = {
        "D2 v3/D2s v3": "Standard_D2s_v3",
        "D4 v3/D4s v3": "Standard_D4s_v3",
        "F4s v2": "Standard_F4s_v2",
        "F8s v2": "Standard_F8s_v2",
        "B2s": "Standard_B2s",
        "DS2 v2": "Standard_DS2_v2",
    }
    for meter, vm_size in known_meter_map.items():
        if meter in text:
            return vm_size
    return ""


def load_reservation_source(source_workbook: Path) -> dict[str, dict[str, Any]]:
    wb = load_workbook(source_workbook, data_only=True, read_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active
    idx = load_headers(ws)

    required = ["Reservation ID", "SKU Name", "Product Name", "Quantity", "Term and Billing Cycle", "Subscription End Date"]
    missing = [header for header in required if header not in idx]
    if missing:
        raise ValueError(f"Reservation source workbook is missing required headers: {missing}")

    reservations: dict[str, dict[str, Any]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        reservation_id = normalize(row[idx["Reservation ID"]])
        if not reservation_id:
            continue
        sku_name = normalize(row[idx["SKU Name"]])
        product_name = normalize(row[idx["Product Name"]])
        vm_size = vm_size_from_text(sku_name, product_name)
        if not vm_size:
            continue
        end_date = parse_date(row[idx["Subscription End Date"]])
        if end_date is None:
            continue
        reservations[reservation_id] = {
            "reservation_id": reservation_id,
            "sku_name": sku_name,
            "product_name": product_name,
            "vm_size": vm_size,
            "quantity": to_number(row[idx["Quantity"]]),
            "term": normalize(row[idx["Term and Billing Cycle"]]),
            "end_date": end_date,
        }
    return reservations


def load_payg_prices(price_workbook: Path) -> dict[tuple[str, str], float]:
    wb = load_workbook(price_workbook, data_only=True, read_only=True)
    ws = wb["RI PAYG Price"] if "RI PAYG Price" in wb.sheetnames else wb.active
    idx = load_headers(ws)

    required = ["VM Size", "Region", "PAYG Monthly Price"]
    missing = [header for header in required if header not in idx]
    if missing:
        raise ValueError(f"PAYG price workbook is missing required headers: {missing}")

    prices: dict[tuple[str, str], float] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        vm_size = normalize(row[idx["VM Size"]])
        region = normalize(row[idx["Region"]])
        price = to_number(row[idx["PAYG Monthly Price"]])
        if vm_size and region and price > 0:
            prices[(vm_size, region.lower())] = price
    return prices


def collect_current_payg(draft_workbooks: list[Path]) -> dict[tuple[str, str], float]:
    payg_by_aks_vm: dict[tuple[str, str], float] = defaultdict(float)
    for workbook_path in draft_workbooks:
        wb = load_workbook(workbook_path, data_only=True, read_only=True)
        for ws in wb.worksheets:
            if not ws.title.startswith("Detail_"):
                continue
            idx = load_headers(ws)
            required = ["Resource Group Name", "Category", "Meter Name", "Product Name", "Tags_JSON", "Total"]
            if any(header not in idx for header in required):
                continue
            for row in ws.iter_rows(min_row=2, values_only=True):
                resource_group = normalize(row[idx["Resource Group Name"]])
                if not resource_group.upper().startswith("MC_"):
                    continue
                if normalize(row[idx["Tags_JSON"]]) == "Reserved Instance":
                    continue
                if normalize(row[idx["Category"]]) != "Virtual Machines":
                    continue
                vm_size = vm_size_from_text(row[idx["Meter Name"]], row[idx["Product Name"]])
                if not vm_size:
                    continue
                payg_by_aks_vm[(aks_name_from_rg(resource_group), vm_size)] += to_number(row[idx["Total"]])
    return payg_by_aks_vm


def collect_reserved_usage(
    reserved_report_workbook: Path,
    reservations: dict[str, dict[str, Any]],
    expiry_start: date | None,
    expiry_end: date | None,
) -> tuple[
    dict[tuple[str, str], float],
    dict[tuple[str, str], float],
    dict[tuple[date, str, str], float],
    dict[tuple[date, str, str], float],
    dict[tuple[date, str, str], set[str]],
    dict[tuple[date, str, str], set[str]],
    dict[tuple[date, str, str], set[str]],
    dict[tuple[str, str], set[str]],
    dict[tuple[str, str], set[str]],
]:
    wb = load_workbook(reserved_report_workbook, data_only=True, read_only=True)
    ws = wb["ReservedInstance"]
    idx = load_headers(ws)

    required = ["Product ID", "Resource Group Name", "Resource Name", "Quantity", "Total", "MD Name"]
    missing = [header for header in required if header not in idx]
    if missing:
        raise ValueError(f"Reserved report workbook is missing required headers: {missing}")

    current_cost: dict[tuple[str, str], float] = defaultdict(float)
    current_qty: dict[tuple[str, str], float] = defaultdict(float)
    expiring_cost: dict[tuple[date, str, str], float] = defaultdict(float)
    expiring_qty: dict[tuple[date, str, str], float] = defaultdict(float)
    expiring_md_names: dict[tuple[date, str, str], set[str]] = defaultdict(set)
    expiring_resource_names: dict[tuple[date, str, str], set[str]] = defaultdict(set)
    expiring_reservation_ids: dict[tuple[date, str, str], set[str]] = defaultdict(set)
    md_names: dict[tuple[str, str], set[str]] = defaultdict(set)
    resource_names: dict[tuple[str, str], set[str]] = defaultdict(set)

    for row in ws.iter_rows(min_row=2, values_only=True):
        reservation_id = normalize(row[idx["Product ID"]])
        reservation = reservations.get(reservation_id)
        if not reservation:
            continue
        resource_group = normalize(row[idx["Resource Group Name"]])
        if not resource_group.upper().startswith("MC_"):
            continue

        end_date = reservation["end_date"]
        vm_size = reservation["vm_size"]
        key = (aks_name_from_rg(resource_group), vm_size)
        cost = to_number(row[idx["Total"]])
        quantity = to_number(row[idx["Quantity"]])

        current_cost[key] += cost
        current_qty[key] += quantity
        md_name = normalize(row[idx["MD Name"]])
        resource_name = normalize(row[idx["Resource Name"]])
        if md_name:
            md_names[key].add(md_name)
        if resource_name:
            resource_names[key].add(resource_name)

        if expiry_start and end_date < expiry_start:
            continue
        if expiry_end and end_date > expiry_end:
            continue

        expiring_key = (end_date, key[0], key[1])
        expiring_cost[expiring_key] += cost
        expiring_qty[expiring_key] += quantity
        if md_name:
            expiring_md_names[expiring_key].add(md_name)
        if resource_name:
            expiring_resource_names[expiring_key].add(resource_name)
        expiring_reservation_ids[expiring_key].add(reservation_id)

    return (
        current_cost,
        current_qty,
        expiring_cost,
        expiring_qty,
        expiring_md_names,
        expiring_resource_names,
        expiring_reservation_ids,
        md_names,
        resource_names,
    )


def build_rows(args: argparse.Namespace) -> list[dict[str, Any]]:
    reservations = load_reservation_source(args.reservation_source)
    prices = load_payg_prices(args.payg_price_workbook)
    payg_by_aks_vm = collect_current_payg(args.draft_workbook)
    expiry_start = parse_date(args.expiry_start) if args.expiry_start else None
    expiry_end = parse_date(args.expiry_end) if args.expiry_end else None
    (
        current_cost,
        current_qty,
        expiring_cost,
        expiring_qty,
        expiring_md_names,
        expiring_resource_names,
        expiring_reservation_ids,
        md_names,
        resource_names,
    ) = collect_reserved_usage(args.reserved_report_workbook, reservations, expiry_start, expiry_end)

    rows: list[dict[str, Any]] = []
    for expiring_key, expired_cost in expiring_cost.items():
        if expired_cost <= 0:
            continue
        end_date, aks_name, vm_size = expiring_key
        key = (aks_name, vm_size)
        current_ri_cost = current_cost[key]
        current_payg = payg_by_aks_vm.get(key, 0.0)
        workload_cost = current_ri_cost + current_payg
        ri_after_expiry = max(current_ri_cost - expired_cost, 0.0)
        current_coverage = (current_ri_cost / workload_cost * 100) if workload_cost else 100.0
        after_coverage = (ri_after_expiry / workload_cost * 100) if workload_cost else 100.0
        payg_price = prices.get((vm_size, args.region.lower()), 0.0)
        rows.append(
            {
                "expiry_date": end_date.isoformat(),
                "reservation_id": ", ".join(sorted(expiring_reservation_ids[expiring_key])),
                "aks_name": aks_name,
                "md": ", ".join(sorted(expiring_md_names[expiring_key] or md_names[key])),
                "vm_size": vm_size,
                "expiring_qty": expiring_qty[expiring_key],
                "payg_price": payg_price,
                "current_payg": current_payg,
                "current_coverage": current_coverage,
                "after_coverage": after_coverage,
                "current_reservation_total": expired_cost,
                "forecast_payg_risk": current_payg + (expiring_qty[expiring_key] * payg_price),
                "resource_names": ", ".join(sorted(expiring_resource_names[expiring_key] or resource_names[key])),
            }
        )

    rows.sort(key=lambda row: (row["expiry_date"], row["aks_name"], row["vm_size"]))
    return rows


def markdown_table(rows: list[dict[str, Any]]) -> str:
    lines = [
        "| Expiry Date | Reservation ID | AKS | MD | VM Size | Expiring RI Qty | PAYG Monthly Price | Current PAYG Left | Current Coverage | After Expired RI Coverage | Current Reservation Total | Forecast PAYG Risk | Used By Resource Name |",
        "| --- | --- | --- | --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | --- |",
    ]
    for row in rows:
        lines.append(
            "| "
            + " | ".join(
                [
                    row["expiry_date"],
                    row["reservation_id"],
                    row["aks_name"],
                    row["md"],
                    row["vm_size"],
                    decimal_comma(row["expiring_qty"]),
                    rupiah(row["payg_price"]) if row["payg_price"] else "Need Price Input",
                    rupiah(row["current_payg"]),
                    percent(row["current_coverage"]),
                    percent(row["after_coverage"]),
                    rupiah(row["current_reservation_total"]),
                    rupiah(row["forecast_payg_risk"]) if row["payg_price"] else "Need Price Input",
                    row["resource_names"],
                ]
            )
            + " |"
        )
    return "\n".join(lines) + "\n"


def tsv_table(rows: list[dict[str, Any]]) -> str:
    headers = [
        "Expiry Date",
        "Reservation ID",
        "AKS",
        "MD",
        "VM Size",
        "Expiring RI Qty",
        "PAYG Monthly Price",
        "Current PAYG Left",
        "Current Coverage",
        "After Expired RI Coverage",
        "Current Reservation Total",
        "Forecast PAYG Risk",
        "Used By Resource Name",
    ]
    lines = ["\t".join(headers)]
    for row in rows:
        lines.append(
            "\t".join(
                [
                    row["expiry_date"],
                    row["reservation_id"],
                    row["aks_name"],
                    row["md"],
                    row["vm_size"],
                    decimal_comma(row["expiring_qty"]),
                    rupiah(row["payg_price"]) if row["payg_price"] else "Need Price Input",
                    rupiah(row["current_payg"]),
                    percent(row["current_coverage"]),
                    percent(row["after_coverage"]),
                    rupiah(row["current_reservation_total"]),
                    rupiah(row["forecast_payg_risk"]) if row["payg_price"] else "Need Price Input",
                    row["resource_names"],
                ]
            )
        )
    return "\n".join(lines) + "\n"


EXCEL_HEADERS = [
    "Expiry Date",
    "Reservation ID",
    "AKS",
    "MD",
    "VM Size",
    "Expiring RI Qty",
    "PAYG Monthly Price",
    "Current PAYG Left",
    "Current Coverage",
    "After Expired RI Coverage",
    "Current Reservation Total",
    "Forecast PAYG Risk",
    "Used By Resource Name",
]


def write_excel(rows: list[dict[str, Any]], workbook_path: Path, sheet_name: str) -> None:
    wb = load_workbook(workbook_path)
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet(sheet_name)

    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(
            [
                row["expiry_date"],
                row["reservation_id"],
                row["aks_name"],
                row["md"],
                row["vm_size"],
                row["expiring_qty"],
                row["payg_price"] if row["payg_price"] else None,
                row["current_payg"],
                row["current_coverage"] / 100,
                row["after_coverage"] / 100,
                row["current_reservation_total"],
                row["forecast_payg_risk"] if row["payg_price"] else None,
                row["resource_names"],
            ]
        )

    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row_idx, 6).number_format = "0.0"
        for col_idx in (7, 8, 11, 12):
            ws.cell(row_idx, col_idx).number_format = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'
        for col_idx in (9, 10):
            ws.cell(row_idx, col_idx).number_format = "0.00%"

    widths = {
        "A": 24,
        "B": 45,
        "C": 18,
        "D": 14,
        "E": 20,
        "F": 16,
        "G": 20,
        "H": 18,
        "I": 18,
        "J": 26,
        "K": 20,
        "L": 20,
        "M": 60,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    wb.save(workbook_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Forecast AKS RI expiry impact with PAYG price input.")
    parser.add_argument("--reservation-source", required=True, type=Path, help="Original reserved-instance workbook.")
    parser.add_argument("--reserved-report-workbook", required=True, type=Path, help="Clean reserved-instance report workbook.")
    parser.add_argument("--payg-price-workbook", required=True, type=Path, help="PAYG price workbook.")
    parser.add_argument("--draft-workbook", required=True, type=Path, action="append", help="Draft workbook. Repeat for 4MD and 9MD.")
    parser.add_argument("--expiry-start", help="Optional expiry start date, YYYY-MM-DD.")
    parser.add_argument("--expiry-end", help="Optional expiry end date, YYYY-MM-DD.")
    parser.add_argument("--region", default="Southeast Asia", help="PAYG price lookup region.")
    parser.add_argument("--output", type=Path, help="Optional markdown output path.")
    parser.add_argument("--format", choices=["markdown", "tsv"], default="markdown", help="Output format.")
    parser.add_argument("--excel-output", type=Path, help="Optional Excel workbook output path.")
    parser.add_argument("--excel-sheet", default="Forecast RI Expiry", help="Excel output sheet name.")
    args = parser.parse_args()

    rows = build_rows(args)
    output_text = markdown_table(rows) if args.format == "markdown" else tsv_table(rows)

    if args.excel_output:
        write_excel(rows, args.excel_output, args.excel_sheet)
        print(f"[DONE] Wrote {len(rows)} row(s) to {args.excel_output} sheet '{args.excel_sheet}'")

    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(output_text, encoding="utf-8")
        print(f"[DONE] Wrote {len(rows)} row(s) to {args.output}")
    elif not args.excel_output:
        print(output_text, end="")


if __name__ == "__main__":
    main()
