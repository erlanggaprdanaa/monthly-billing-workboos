from __future__ import annotations

import argparse
import shutil
import tempfile
from pathlib import Path
from typing import Any

import win32com.client as win32
from openpyxl import load_workbook


SOURCE_HEADERS = [
    "Period",
    "Agreement Name",
    "Company Name",
    "Tenant Primary Domain",
    "Publisher Subscription ID",
    "Publisher Subscription Name",
    "Product ID",
    "Product Name",
    "Category",
    "Subcategory",
    "Meter Name",
    "Usage Start Date",
    "Usage End Date",
    "Quantity",
    "Unit",
    "Currency",
    "Resource Group Name",
    "Resource Name",
    "Location",
    "Tags_JSON",
    "Tags_KeyValue",
    "Tags_Count",
    "Total",
    "MD Name",
]

DETAIL_HEADERS = SOURCE_HEADERS[:-1]

DETAIL_RULES = {
    "Anper": "Detail_Anper",
    "Sinsen": "Detail_Sinsen",
    "CDN": "Detail_CDN",
    "MenaraAgung": "Detail_MenaraAgung",
    "MSK": "Detail_MSK",
    "HPM": "Detail_HPM",
    "ITC": "Detail_ITC",
    "TDM": "Detail_TDM",
    "TRIO": "Detail_TRIO",
    "Zenobia": "Detail_Zenobia",
}

ACCOUNTING_IDR_FORMAT = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'
ACCOUNTING_IDR_FORMAT_LOCAL = '_-[$Rp-421] * #.##0_-;\\-[$Rp-421] * #.##0_-;_-[$Rp-421] * "-"_-;_-@_-'
QUANTITY_NUMBER_FORMAT = "General"
XL_UP = -4162
XL_TO_LEFT = -4159


def normalize_header(value: object) -> str:
    return "" if value is None else str(value).strip()


def load_source_workbook_safe(path: Path):
    try:
        return load_workbook(path, read_only=True, keep_links=False, data_only=False)
    except PermissionError:
        temp_copy = Path(tempfile.gettempdir()) / f"{path.stem}-copy{path.suffix}"
        shutil.copy2(path, temp_copy)
        return load_workbook(temp_copy, read_only=True, keep_links=False, data_only=False)


def get_openpyxl_header_map(sheet) -> dict[str, int]:
    headers: dict[str, int] = {}
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    for col_idx, value in enumerate(first_row, start=1):
        header = normalize_header(value)
        if header:
            headers[header] = col_idx
    return headers


def ensure_headers(header_map: dict[str, int], required_headers: list[str], sheet_name: str) -> None:
    missing = [header for header in required_headers if header not in header_map]
    if missing:
        raise ValueError(f"Sheet '{sheet_name}' is missing required headers: {missing}")


def collect_rows_by_md(source_sheet, source_header_map: dict[str, int]) -> dict[str, list[list[Any]]]:
    rows_by_md: dict[str, list[list[Any]]] = {}

    for row in source_sheet.iter_rows(min_row=2, values_only=True):
        row_values = [row[source_header_map[header] - 1] for header in SOURCE_HEADERS]
        if not any(value is not None and value != "" for value in row_values):
            continue

        md_name = row[source_header_map["MD Name"] - 1]
        if md_name is None or str(md_name).strip() == "":
            continue

        rows_by_md.setdefault(str(md_name).strip(), []).append(row_values)

    return rows_by_md


def get_excel_header_map(sheet) -> dict[str, int]:
    last_col = sheet.Cells(1, sheet.Columns.Count).End(XL_TO_LEFT).Column
    values = sheet.Range(sheet.Cells(1, 1), sheet.Cells(1, last_col)).Value
    header_map: dict[str, int] = {}
    for col_idx in range(1, last_col + 1):
        value = values[0][col_idx - 1] if isinstance(values, tuple) else values
        header = normalize_header(value)
        if header:
            header_map[header] = col_idx
    return header_map


def worksheet_exists(workbook, name: str) -> bool:
    try:
        workbook.Worksheets(name)
        return True
    except Exception:
        return False


def find_next_empty_row(sheet) -> int:
    last_row = sheet.Cells(sheet.Rows.Count, 1).End(XL_UP).Row
    return last_row + 1


def delete_existing_reserved_instance_rows(sheet) -> int:
    header_map = get_excel_header_map(sheet)
    ensure_headers(header_map, DETAIL_HEADERS, sheet.Name)

    tags_json_col = header_map["Tags_JSON"]
    deleted = 0
    last_row = sheet.Cells(sheet.Rows.Count, 1).End(XL_UP).Row
    for row_idx in range(last_row, 1, -1):
        tags_json = normalize_header(sheet.Cells(row_idx, tags_json_col).Value)
        if tags_json == "Reserved Instance":
            sheet.Rows(row_idx).Delete()
            deleted += 1
    return deleted


def apply_formats(sheet, start_row: int, end_row: int) -> None:
    if end_row < start_row:
        return
    sheet.Range(sheet.Cells(start_row, 14), sheet.Cells(end_row, 14)).NumberFormat = QUANTITY_NUMBER_FORMAT
    sheet.Range(sheet.Cells(start_row, 15), sheet.Cells(end_row, 15)).NumberFormat = ACCOUNTING_IDR_FORMAT
    sheet.Range(sheet.Cells(start_row, 23), sheet.Cells(end_row, 23)).NumberFormat = ACCOUNTING_IDR_FORMAT
    try:
        sheet.Range(sheet.Cells(start_row, 14), sheet.Cells(end_row, 14)).NumberFormatLocal = QUANTITY_NUMBER_FORMAT
        sheet.Range(sheet.Cells(start_row, 15), sheet.Cells(end_row, 15)).NumberFormatLocal = ACCOUNTING_IDR_FORMAT_LOCAL
        sheet.Range(sheet.Cells(start_row, 23), sheet.Cells(end_row, 23)).NumberFormatLocal = ACCOUNTING_IDR_FORMAT_LOCAL
    except Exception:
        pass


def append_rows_to_detail_sheet(sheet, rows: list[list[Any]]) -> tuple[int, int]:
    header_map = get_excel_header_map(sheet)
    ensure_headers(header_map, DETAIL_HEADERS, sheet.Name)

    start_row = find_next_empty_row(sheet)
    for header in DETAIL_HEADERS:
        col_idx = header_map[header]
        values = [[row[SOURCE_HEADERS.index(header)]] for row in rows]
        end_row = start_row + len(rows) - 1
        sheet.Range(sheet.Cells(start_row, col_idx), sheet.Cells(end_row, col_idx)).Value = tuple(tuple(item) for item in values)

    if rows:
        apply_formats(sheet, start_row, start_row + len(rows) - 1)

    return start_row, len(rows)


def process(source_workbook_path: Path, target_workbook_path: Path, source_sheet_name: str) -> list[str]:
    print("[INFO] Loading source workbook")
    source_wb = load_source_workbook_safe(source_workbook_path)
    if source_sheet_name not in source_wb.sheetnames:
        raise ValueError(f"Source sheet '{source_sheet_name}' not found.")

    source_sheet = source_wb[source_sheet_name]
    source_header_map = get_openpyxl_header_map(source_sheet)
    ensure_headers(source_header_map, SOURCE_HEADERS, source_sheet.title)

    print("[INFO] Collecting rows by MD Name")
    rows_by_md = collect_rows_by_md(source_sheet, source_header_map)
    print(f"[INFO] Found MD groups: {', '.join(sorted(rows_by_md))}")

    print("[INFO] Opening target workbook in Excel")
    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False

    workbook = None
    results: list[str] = []
    try:
        workbook = excel.Workbooks.Open(str(target_workbook_path))

        for md_name, target_sheet_name in DETAIL_RULES.items():
            if not worksheet_exists(workbook, target_sheet_name):
                raise ValueError(f"Target sheet '{target_sheet_name}' not found in workbook.")

            target_sheet = workbook.Worksheets(target_sheet_name)
            deleted = delete_existing_reserved_instance_rows(target_sheet)
            if deleted:
                results.append(f"[CLEAN] {target_sheet_name}: deleted {deleted} existing Reserved Instance row(s)")

            rows = rows_by_md.get(md_name, [])
            if not rows:
                results.append(f"[SKIP] {md_name}: no rows found in source workbook")
                continue
            print(f"[INFO] Appending {len(rows)} row(s) into '{target_sheet_name}'")
            start_row, copied = append_rows_to_detail_sheet(target_sheet, rows)
            results.append(f"[APPEND] {target_sheet_name}: copied {copied} row(s) starting at row {start_row}")

        print("[INFO] Saving target workbook")
        workbook.Save()
        return results
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=True)
        excel.Quit()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Append Reserved Instance rows by MD Name into Draft-9MD workbook detail sheets."
    )
    parser.add_argument("--source-workbook", required=True, type=Path, help="Path to reserved instance report workbook.")
    parser.add_argument("--target-workbook", required=True, type=Path, help="Path to Draft-9MD workbook.")
    parser.add_argument("--source-sheet", default="ReservedInstance", help="Source sheet name.")
    args = parser.parse_args()

    print(f"[INFO] Source workbook : {args.source_workbook}")
    print(f"[INFO] Target workbook : {args.target_workbook}")
    print(f"[INFO] Source sheet    : {args.source_sheet}")

    results = process(args.source_workbook, args.target_workbook, args.source_sheet)
    for line in results:
        print(line)
    print("[DONE] Reserved Instance append to Draft-9MD completed.")


if __name__ == "__main__":
    main()
