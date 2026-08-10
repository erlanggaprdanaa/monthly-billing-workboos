from __future__ import annotations

import argparse
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import re

from openpyxl import load_workbook

RI_PATTERN = re.compile(r"\bRI\b|Reserved Instance|Reservation", re.IGNORECASE)


def round_idr(value: Decimal) -> int:
    return int(value.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def pct_change(previous: Decimal, current: Decimal) -> str:
    if previous == 0:
        if current == 0:
            return "0.00%"
        return "new"
    pct = ((current - previous) / previous) * Decimal("100")
    return f"{pct.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)}%"


def is_reserved_instance_row(row_map: dict[str, object]) -> bool:
    text = " ".join(
        str(row_map.get(key) or "")
        for key in ["Product Name", "Category", "Subcategory", "Meter Name", "Tags_JSON", "Tags_KeyValue"]
    )
    return bool(RI_PATTERN.search(text))


def is_google_secops_row(row_map: dict[str, object]) -> bool:
    resource_group = str(row_map.get("Resource Group Name") or "").lower()
    resource_name = str(row_map.get("Resource Name") or "").lower()
    product_name = str(row_map.get("Product Name") or "").lower()
    return "secops" in resource_group or "secops" in resource_name or "google secops" in product_name


def build_group_key(row: tuple[object, ...], idx: dict[str, int], group_by: str) -> str:
    if group_by == "product":
        return str(row[idx["Product Name"]] or "(blank)")
    if group_by == "resource":
        return str(row[idx["Resource Name"]] or "(blank)")
    if group_by == "meter":
        return "||".join(
            [
                str(row[idx["Category"]] or ""),
                str(row[idx["Subcategory"]] or ""),
                str(row[idx["Meter Name"]] or ""),
            ]
        )
    raise ValueError(f"Unsupported group_by: {group_by}")


def load_sheet_totals(
    workbook_path: Path,
    sheet_name: str,
    group_by: str,
    *,
    exclude_reserved_instance: bool,
    exclude_secops: bool,
) -> tuple[Decimal, dict[str, Decimal]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True, keep_links=False)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{sheet_name}' not found in {workbook_path}")

    ws = wb[sheet_name]
    headers = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    idx = {str(h): i for i, h in enumerate(headers)}

    required = ["Total", "Product Name", "Category", "Subcategory", "Meter Name", "Resource Name"]
    missing = [name for name in required if name not in idx]
    if missing:
        raise ValueError(f"Missing columns in {workbook_path} / {sheet_name}: {', '.join(missing)}")

    total_all = Decimal("0")
    grouped: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))

    for row in ws.iter_rows(min_row=2, values_only=True):
        total_value = row[idx["Total"]]
        if total_value in (None, ""):
            continue

        row_map = {str(h): row[i] for i, h in enumerate(headers)}
        if exclude_reserved_instance and is_reserved_instance_row(row_map):
            continue
        if exclude_secops and is_google_secops_row(row_map):
            continue

        total_decimal = Decimal(str(total_value))
        total_all += total_decimal

        key = build_group_key(row, idx, group_by)
        grouped[key] += total_decimal

    return total_all, grouped


def format_idr(value: Decimal | int) -> str:
    if not isinstance(value, Decimal):
        value = Decimal(value)
    return f"Rp {round_idr(value):,}".replace(",", ".")


def md_escape(value: str) -> str:
    return value.replace("|", "\\|")


def render_service_cells(group_by: str, key: str) -> list[str]:
    if group_by != "meter":
        return [md_escape(key)]

    category, subcategory, meter_name = (key.split("||", 2) + ["", ""])[:3]
    detail = meter_name or subcategory
    return [md_escape(category or "(blank)"), md_escape(detail or "(blank)")]


def safe_name(value: str) -> str:
    sanitized = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip())
    sanitized = re.sub(r"-{2,}", "-", sanitized).strip("-")
    return sanitized or "report"


def infer_period_label(workbook_path: Path) -> str:
    parent_name = workbook_path.parent.name.strip()
    return parent_name or workbook_path.stem


def build_markdown(
    previous: Path,
    current: Path,
    sheet: str,
    group_by: str,
    prev_total: Decimal,
    curr_total: Decimal,
    direction: str,
    increases: list[tuple[str, Decimal, Decimal, Decimal]],
    decreases: list[tuple[str, Decimal, Decimal, Decimal]],
    top: int,
) -> str:
    lines: list[str] = []
    lines.append(f"# Compare {sheet}")
    lines.append("")
    lines.append(f"- Previous workbook: `{previous}`")
    lines.append(f"- Current workbook: `{current}`")
    lines.append(f"- Group by: `{group_by}`")
    lines.append("")
    lines.append("## Summary")
    lines.append("")
    lines.append("| Metric | Value |")
    lines.append("|---|---:|")
    lines.append(f"| Previous total | {format_idr(prev_total)} |")
    lines.append(f"| Current total | {format_idr(curr_total)} |")
    lines.append(f"| Delta | {format_idr((curr_total - prev_total).copy_abs())} ({direction}) |")
    lines.append(f"| Change | {pct_change(prev_total, curr_total)} |")
    lines.append("")
    lines.append(f"## Top {top} Increases")
    lines.append("")
    if not increases:
        lines.append("- None")
    else:
        total_increase = sum((item[3] for item in increases), Decimal("0"))
        lines.append(f"**Total Increase: +{format_idr(total_increase)}**")
        lines.append("")
        if group_by == "meter":
            lines.append("| No | Category | Service | Previous | Current | Delta | Change |")
            lines.append("|---:|---|---|---:|---:|---:|---:|")
        else:
            lines.append("| No | Service | Previous | Current | Delta | Change |")
            lines.append("|---:|---|---:|---:|---:|---:|")
        for idx, (key, prev_value, curr_value, delta) in enumerate(increases, start=1):
            service_cells = render_service_cells(group_by, key)
            if group_by == "meter":
                lines.append(
                    f"| {idx} | {service_cells[0]} | {service_cells[1]} | "
                    f"{format_idr(prev_value)} | {format_idr(curr_value)} | "
                    f"+{format_idr(delta)} | {pct_change(prev_value, curr_value)} |"
                )
            else:
                lines.append(
                    f"| {idx} | {service_cells[0]} | {format_idr(prev_value)} | {format_idr(curr_value)} | "
                    f"+{format_idr(delta)} | {pct_change(prev_value, curr_value)} |"
                )
    lines.append("")
    lines.append(f"## Top {top} Decreases")
    lines.append("")
    if not decreases:
        lines.append("- None")
    else:
        total_decrease = sum((item[3] for item in decreases), Decimal("0"))
        lines.append(f"**Total Decrease: -{format_idr(total_decrease.copy_abs())}**")
        lines.append("")
        if group_by == "meter":
            lines.append("| No | Category | Service | Previous | Current | Delta | Change |")
            lines.append("|---:|---|---|---:|---:|---:|---:|")
        else:
            lines.append("| No | Service | Previous | Current | Delta | Change |")
            lines.append("|---:|---|---:|---:|---:|---:|")
        for idx, (key, prev_value, curr_value, delta) in enumerate(decreases, start=1):
            service_cells = render_service_cells(group_by, key)
            if group_by == "meter":
                lines.append(
                    f"| {idx} | {service_cells[0]} | {service_cells[1]} | "
                    f"{format_idr(prev_value)} | {format_idr(curr_value)} | "
                    f"-{format_idr(delta.copy_abs())} | {pct_change(prev_value, curr_value)} |"
                )
            else:
                lines.append(
                    f"| {idx} | {service_cells[0]} | {format_idr(prev_value)} | {format_idr(curr_value)} | "
                    f"-{format_idr(delta.copy_abs())} | {pct_change(prev_value, curr_value)} |"
                )
    lines.append("")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare monthly totals for a Detail_* sheet.")
    parser.add_argument("--previous", required=True, type=Path, help="Previous month workbook")
    parser.add_argument("--current", required=True, type=Path, help="Current month workbook")
    parser.add_argument("--sheet", required=True, help="Sheet name to compare, e.g. Detail_Astra")
    parser.add_argument(
        "--group-by",
        choices=["product", "meter", "resource"],
        default="product",
        help="Grouping key for top changes",
    )
    parser.add_argument("--top", type=int, default=10, help="Top N increases to show")
    parser.add_argument(
        "--exclude-reserved-instance",
        action="store_true",
        help="Exclude rows detected as Reserved Instance from totals and comparisons",
    )
    parser.add_argument(
        "--exclude-secops",
        action="store_true",
        help="Exclude rows detected as Google SecOps from totals and comparisons",
    )
    parser.add_argument("--output-dir", type=Path, help="Optional output folder for markdown report")
    parser.add_argument("--output-md", type=Path, help="Optional explicit markdown output file path")
    args = parser.parse_args()

    prev_total, prev_grouped = load_sheet_totals(
        args.previous,
        args.sheet,
        args.group_by,
        exclude_reserved_instance=args.exclude_reserved_instance,
        exclude_secops=args.exclude_secops,
    )
    curr_total, curr_grouped = load_sheet_totals(
        args.current,
        args.sheet,
        args.group_by,
        exclude_reserved_instance=args.exclude_reserved_instance,
        exclude_secops=args.exclude_secops,
    )

    delta_total = curr_total - prev_total
    direction = "Naik" if delta_total > 0 else "Turun" if delta_total < 0 else "Flat"

    changes = []
    all_keys = set(prev_grouped) | set(curr_grouped)
    for key in all_keys:
        prev_value = prev_grouped.get(key, Decimal("0"))
        curr_value = curr_grouped.get(key, Decimal("0"))
        delta = curr_value - prev_value
        changes.append((key, prev_value, curr_value, delta))

    increases = sorted(
        [item for item in changes if item[3] > 0],
        key=lambda item: item[3],
        reverse=True,
    )[: args.top]

    decreases = sorted(
        [item for item in changes if item[3] < 0],
        key=lambda item: item[3],
    )[: args.top]

    output_md = args.output_md
    if output_md is None and args.output_dir is not None:
        previous_label = safe_name(infer_period_label(args.previous))
        current_label = safe_name(infer_period_label(args.current))
        sheet_label = safe_name(args.sheet)
        output_md = args.output_dir / f"{sheet_label}-{previous_label}-vs-{current_label}.md"

    if output_md is not None:
        output_md.parent.mkdir(parents=True, exist_ok=True)
        markdown = build_markdown(
            previous=args.previous,
            current=args.current,
            sheet=args.sheet,
            group_by=args.group_by,
            prev_total=prev_total,
            curr_total=curr_total,
            direction=direction,
            increases=increases,
            decreases=decreases,
            top=args.top,
        )
        output_md.write_text(markdown, encoding="utf-8")
        print(f"Markdown report saved: {output_md}")
    else:
        print(build_markdown(
            previous=args.previous,
            current=args.current,
            sheet=args.sheet,
            group_by=args.group_by,
            prev_total=prev_total,
            curr_total=curr_total,
            direction=direction,
            increases=increases,
            decreases=decreases,
            top=args.top,
        ))


if __name__ == "__main__":
    main()
