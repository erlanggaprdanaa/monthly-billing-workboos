from __future__ import annotations

import argparse
from pathlib import Path

import win32com.client as win32
from openpyxl import load_workbook


EXPECTED_HEADERS = [
    "Period",
    "Agreement Name",
    "Company Name",
    "Tenant Primary Domain",
    "Publisher Subscription ID",
    "Publisher Subscription Name",
    "Product ID",
    "Product Name",
    "Category",
    "Subcategory",
    "Meter Name",
    "Usage Start Date",
    "Usage End Date",
    "Quantity",
    "Unit",
    "Currency",
    "Resource Group Name",
    "Resource Name",
    "Location",
    "Tags_JSON",
    "Tags_KeyValue",
    "Tags_Count",
    "Total",
    "MD Name",
]

ACCOUNTING_IDR_FORMAT = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'
ACCOUNTING_IDR_FORMAT_LOCAL = '_-[$Rp-421] * #.##0_-;\\-[$Rp-421] * #.##0_-;_-[$Rp-421] * "-"_-;_-@_-'
QUANTITY_NUMBER_FORMAT = '[=0]0.00;[<1]0.00;0.0#'
XL_UP = -4162
XL_TO_LEFT = -4159


def normalize(value: object) -> str:
    return "" if value is None else str(value).strip()


def read_template_rows(path: Path, sheet_name: str) -> list[list[object]]:
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    ws = wb[sheet_name]
    headers = [normalize(c) for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    if headers != EXPECTED_HEADERS:
        raise ValueError("Template headers do not match expected Reserved Instance headers.")
    rows: list[list[object]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None and v != "" for v in row):
            rows.append(list(row))
    return rows


def get_header_map(sheet) -> dict[str, int]:
    last_col = sheet.Cells(1, sheet.Columns.Count).End(XL_TO_LEFT).Column
    values = sheet.Range(sheet.Cells(1, 1), sheet.Cells(1, last_col)).Value
    header_map: dict[str, int] = {}
    for col_idx in range(1, last_col + 1):
        value = values[0][col_idx - 1] if isinstance(values, tuple) else values
        header = normalize(value)
        if header:
            header_map[header] = col_idx
    return header_map


def apply_formats(sheet, start_row: int, end_row: int) -> None:
    if end_row < start_row:
        return
    sheet.Range(sheet.Cells(start_row, 14), sheet.Cells(end_row, 14)).NumberFormat = QUANTITY_NUMBER_FORMAT
    sheet.Range(sheet.Cells(start_row, 15), sheet.Cells(end_row, 15)).NumberFormat = ACCOUNTING_IDR_FORMAT
    sheet.Range(sheet.Cells(start_row, 23), sheet.Cells(end_row, 23)).NumberFormat = ACCOUNTING_IDR_FORMAT
    try:
        sheet.Range(sheet.Cells(start_row, 15), sheet.Cells(end_row, 15)).NumberFormatLocal = ACCOUNTING_IDR_FORMAT_LOCAL
        sheet.Range(sheet.Cells(start_row, 23), sheet.Cells(end_row, 23)).NumberFormatLocal = ACCOUNTING_IDR_FORMAT_LOCAL
    except Exception:
        pass


def main() -> None:
    parser = argparse.ArgumentParser(description="Append shared license template rows to Reserved Instance Shared sheet.")
    parser.add_argument("--template-workbook", required=True, type=Path)
    parser.add_argument("--target-workbook", required=True, type=Path)
    parser.add_argument("--template-sheet", default="SharedLicense")
    parser.add_argument("--target-sheet", default="Reserved Instance Shared")
    args = parser.parse_args()

    rows = read_template_rows(args.template_workbook, args.template_sheet)
    if not rows:
        raise ValueError("No template rows found to append.")

    excel = win32.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    workbook = None
    try:
        print(f"[INFO] Opening target workbook : {args.target_workbook}")
        workbook = excel.Workbooks.Open(str(args.target_workbook))
        sheet = workbook.Worksheets(args.target_sheet)
        header_map = get_header_map(sheet)
        missing = [h for h in EXPECTED_HEADERS if h not in header_map]
        if missing:
            raise ValueError(f"Target sheet missing headers: {missing}")

        start_row = sheet.Cells(sheet.Rows.Count, 1).End(XL_UP).Row + 1
        end_row = start_row + len(rows) - 1
        matrix = tuple(tuple(row) for row in rows)
        sheet.Range(sheet.Cells(start_row, 1), sheet.Cells(end_row, len(EXPECTED_HEADERS))).Value = matrix
        apply_formats(sheet, start_row, end_row)
        workbook.Save()
        print(f"[DONE] Appended {len(rows)} row(s) into '{args.target_sheet}' starting at row {start_row}.")
    finally:
        if workbook is not None:
            workbook.Close(SaveChanges=True)
        excel.Quit()


if __name__ == "__main__":
    main()
