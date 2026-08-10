from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


OUTPUT_PATH = Path(r"H:\Work\Code-base\monthly-billing-workbooks\template-shared-license.xlsx")

HEADERS = [
    "Period",
    "Agreement Name",
    "Company Name",
    "Tenant Primary Domain",
    "Publisher Subscription ID",
    "Publisher Subscription Name",
    "Product ID",
    "Product Name",
    "Category",
    "Subcategory",
    "Meter Name",
    "Usage Start Date",
    "Usage End Date",
    "Quantity",
    "Unit",
    "Currency",
    "Resource Group Name",
    "Resource Name",
    "Location",
    "Tags_JSON",
    "Tags_KeyValue",
    "Tags_Count",
    "Total",
    "MD Name",
]

ACCOUNTING_IDR_FORMAT = '_-[$Rp-421]* #,##0_-;\\-[$Rp-421]* #,##0_-;_-[$Rp-421]* "-"_-;_-@_-'


def build_rows() -> list[list[object]]:
    start_date = datetime(2026, 6, 1)
    end_date = datetime(2026, 6, 30)
    common = [
        "June",
        "Microsoft CSP Azure Plan",
        "Divisi Inovasi",
        "divisiinovasi.onmicrosoft.com",
        "a4cdaa2e-caa7-440f-b849-928a3ceb7e31",
        "AzureCSP_DivisiInovasi",
    ]
    tail = [
        "IDR",
        "RG-SHARED",
        None,
        None,
        "Shared License",
        "Shared Allocation",
        1,
        None,
        "Shared",
    ]

    return [
        [*common, "E365", "Microsoft 365 Standard", "License", "Microsoft 365 Standard", "Microsoft 365 Standard", start_date, end_date, 1, 163728, *tail[:2], "microsoft-365-standard", *tail[3:7], 163728, tail[8]],
        [*common, "POWERBI", "Power BI Pro", "Lisensi", "Power BI Pro", "Power BI Pro", start_date, end_date, 2, 165500, *tail[:2], "power-bi-pro", *tail[3:7], 331000, tail[8]],
        [*common, "SENDGRID", "Twilio Sendgrid - Pro 700 k", "Sendgrid", "Twilio Sendgrid - Pro 700 k", "Twilio Sendgrid - Pro 700 k", start_date, end_date, 1, 8516433, *tail[:2], "twilio-sendgrid-pro-700k", *tail[3:7], 8516433, tail[8]],
        [*common, "GITHUB", "Github Lisensi", "Github team license", "Github Lisensi", "Github Lisensi", start_date, end_date, 2, 503219, *tail[:2], "github-team-license", *tail[3:7], 1006438, tail[8]],
    ]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "SharedLicense"
    ws.append(HEADERS)
    for row in build_rows():
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row_idx in range(2, ws.max_row + 1):
        ws.cell(row_idx, 12).number_format = "dd/mm/yyyy"
        ws.cell(row_idx, 13).number_format = "dd/mm/yyyy"
        ws.cell(row_idx, 15).number_format = ACCOUNTING_IDR_FORMAT
        ws.cell(row_idx, 23).number_format = ACCOUNTING_IDR_FORMAT

    widths = {
        "A": 12, "B": 28, "C": 18, "D": 30, "E": 40, "F": 24, "G": 14, "H": 28, "I": 20, "J": 28, "K": 28,
        "L": 14, "M": 14, "N": 10, "O": 14, "P": 10, "Q": 20, "R": 28, "S": 14, "T": 18, "U": 20, "V": 10, "W": 14, "X": 12,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    wb.save(OUTPUT_PATH)
    print(f"Created template: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
